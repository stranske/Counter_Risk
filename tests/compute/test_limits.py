"""Tests for limit breach computations."""

from __future__ import annotations

from pathlib import Path
from typing import Any, cast

import pytest

from counter_risk.compute.limits import (
    _COUNTERPARTY_GRANULARITY,
    _FUTURES_GRANULARITY,
    _LIMIT_GRANULARITY_KEY,
    _find_notional,
    check_limits,
    find_missing_limit_entities,
    write_limit_breaches_csv,
)
from counter_risk.limits_config import LimitsConfig, load_limits_config
from counter_risk.parsers.cprs_fcm import parse_fcm_totals, parse_futures_detail
from counter_risk.pipeline.run import (
    _build_limit_exposure_rows,
    _compute_and_write_concentration_metrics,
)


def _as_records(table: Any) -> list[dict[str, Any]]:
    if hasattr(table, "to_dict"):
        return cast(list[dict[str, Any]], table.to_dict(orient="records"))
    return [dict(row) for row in table]


@pytest.mark.parametrize("blank", [None, "", " \t "])
@pytest.mark.parametrize("fallback", ["Notional", "exposure", "total", "amount"])
def test_find_notional_skips_blank_aliases(blank: Any, fallback: str) -> None:
    assert _find_notional({"notional": blank, fallback: 250000.0}) == 250000.0


@pytest.mark.parametrize("value", [True, False])
@pytest.mark.parametrize("key", ["notional", "Notional", "exposure", "total", "amount"])
def test_find_notional_rejects_boolean_aliases(value: bool, key: str) -> None:
    with pytest.raises(ValueError, match="notional values must be numeric"):
        _find_notional({key: value})


def test_find_notional_preserves_precedence_zero_and_invalid_value_rejection() -> None:
    assert _find_notional({"notional": 0, "exposure": 250000.0}) == 0.0
    assert _find_notional({"notional": "-125.5", "exposure": 250000.0}) == -125.5
    assert _find_notional({"amount": "250000"}) == 250000.0
    for invalid in (True, False, "oops", []):
        with pytest.raises(ValueError, match="notional values must be numeric"):
            _find_notional({"notional": invalid, "exposure": 250000.0})


@pytest.mark.parametrize("row", [{}, {"notional": None, "exposure": "", "amount": " "}])
def test_find_notional_rejects_rows_without_usable_aliases(row: dict[str, Any]) -> None:
    with pytest.raises(ValueError, match="must include one of the notional columns"):
        _find_notional(row)


def test_check_limits_uses_fallback_notional_for_breaches_and_denominator() -> None:
    exposures = [
        {"counterparty": "Alpha", "notional": None, "Notional": "", "exposure": 250000.0},
        {"counterparty": "Beta", "notional": " ", "exposure": None, "amount": 750000.0},
    ]
    config = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": "counterparty",
                "entity_name": "Alpha",
                "limit_kind": kind,
                "limit_value": limit,
            }
            for kind, limit in [("absolute_notional", 200000.0), ("percent_of_total", 0.2)]
        ],
    }
    rows = _as_records(check_limits(exposures, config))
    by_kind = {row["limit_kind"]: row for row in rows}
    assert set(by_kind) == {"absolute_notional", "percent_of_total"}
    assert by_kind["absolute_notional"]["actual_value"] == 250000.0
    assert by_kind["absolute_notional"]["breach_amount"] == 50000.0
    assert by_kind["percent_of_total"]["actual_value"] == pytest.approx(0.25)
    assert by_kind["percent_of_total"]["breach_amount"] == pytest.approx(0.05)


@pytest.mark.parametrize("enabled", [True, False])
def test_finite_limit_preserves_breach_and_disabled_behavior(enabled: bool) -> None:
    config = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": "counterparty",
                "entity_name": "Alpha",
                "limit_kind": "absolute_notional",
                "limit_value": 50,
                "severity": "fail",
                "enabled": enabled,
            }
        ],
    }
    rows = _as_records(check_limits([{"counterparty": "Alpha", "notional": 100}], config))
    if enabled:
        assert len(rows) == 1
        assert rows[0]["actual_value"] == 100
        assert rows[0]["limit_value"] == 50
        assert rows[0]["breach_amount"] == 50
        assert rows[0]["severity"] == "fail"
    else:
        assert rows == []


@pytest.mark.parametrize("limit_kind", ["absolute_notional", "percent_of_total"])
@pytest.mark.parametrize("value", [float("inf"), float("-inf"), float("nan")])
def test_check_limits_rejects_non_finite_policy_before_evaluation(
    monkeypatch: pytest.MonkeyPatch, limit_kind: str, value: float
) -> None:
    def unexpected_evaluation(*args: Any, **kwargs: Any) -> None:
        pytest.fail("Non-finite policy reached exposure evaluation")

    monkeypatch.setattr(
        "counter_risk.compute.limits._denominator_abs_notional", unexpected_evaluation
    )
    config = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": "counterparty",
                "entity_name": "Alpha",
                "limit_kind": limit_kind,
                "limit_value": value,
            }
        ],
    }
    with pytest.raises(ValueError, match=r"limit_value\s+Input should be a finite number"):
        check_limits([{"counterparty": "Alpha", "notional": 100}], config)


def test_check_limits_detects_absolute_and_percent_breaches_across_entity_types() -> None:
    exposures = [
        {
            "counterparty": "Alpha Bank",
            "fcm": "FCM One",
            "clearing_house": "CME",
            "segment": "Treasury",
            "custom_group": "Trend Energy",
            "notional": 90.0,
        },
        {
            "counterparty": "Alpha Bank",
            "fcm": "FCM Two",
            "clearing_house": "ICE",
            "segment": "Equity",
            "custom_group": "Trend Energy",
            "notional": 30.0,
        },
        {
            "counterparty": "Beta Fund",
            "fcm": "FCM One",
            "clearing_house": "CME",
            "segment": "Commodity",
            "custom_group": "Trend Rates",
            "notional": 80.0,
        },
    ]
    limits_cfg = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": "counterparty",
                "entity_name": "Alpha Bank",
                "limit_value": 100.0,
                "limit_kind": "absolute_notional",
                "severity": "fail",
                "notes": "Single-name cap.",
            },
            {
                "entity_type": "fcm",
                "entity_name": "FCM One",
                "limit_value": 0.65,
                "limit_kind": "percent_of_total",
            },
            {
                "entity_type": "clearing_house",
                "entity_name": "CME",
                "limit_value": 0.6,
                "limit_kind": "percent_of_total",
            },
            {
                "entity_type": "segment",
                "entity_name": "Treasury",
                "limit_value": 0.2,
                "limit_kind": "percent_of_total",
            },
            {
                "entity_type": "custom_group",
                "entity_name": "Trend Energy",
                "limit_value": 0.4,
                "limit_kind": "percent_of_total",
            },
        ],
    }

    rows = _as_records(check_limits(exposures, limits_cfg))

    assert rows == [
        {
            "entity_type": "clearing_house",
            "entity_name": "cme",
            "limit_kind": "percent_of_total",
            "severity": "warning",
            "actual_value": pytest.approx(170.0 / 200.0),
            "limit_value": 0.6,
            "breach_amount": pytest.approx((170.0 / 200.0) - 0.6),
            "notes": None,
        },
        {
            "entity_type": "counterparty",
            "entity_name": "alpha_bank",
            "limit_kind": "absolute_notional",
            "severity": "fail",
            "actual_value": 120.0,
            "limit_value": 100.0,
            "breach_amount": 20.0,
            "notes": "Single-name cap.",
        },
        {
            "entity_type": "custom_group",
            "entity_name": "trend_energy",
            "limit_kind": "percent_of_total",
            "severity": "warning",
            "actual_value": pytest.approx(120.0 / 200.0),
            "limit_value": 0.4,
            "breach_amount": pytest.approx((120.0 / 200.0) - 0.4),
            "notes": None,
        },
        {
            "entity_type": "fcm",
            "entity_name": "fcm_one",
            "limit_kind": "percent_of_total",
            "severity": "warning",
            "actual_value": pytest.approx(170.0 / 200.0),
            "limit_value": 0.65,
            "breach_amount": pytest.approx((170.0 / 200.0) - 0.65),
            "notes": None,
        },
        {
            "entity_type": "segment",
            "entity_name": "treasury",
            "limit_kind": "percent_of_total",
            "severity": "warning",
            "actual_value": pytest.approx(90.0 / 200.0),
            "limit_value": 0.2,
            "breach_amount": pytest.approx((90.0 / 200.0) - 0.2),
            "notes": None,
        },
    ]


def test_percent_of_total_scopes_denominator_to_matching_granularity() -> None:
    exposures = [
        {
            _LIMIT_GRANULARITY_KEY: _COUNTERPARTY_GRANULARITY,
            "counterparty": "Alpha",
            "notional": 600.0,
        },
        {
            _LIMIT_GRANULARITY_KEY: _COUNTERPARTY_GRANULARITY,
            "counterparty": "Beta",
            "notional": 400.0,
        },
        {
            _LIMIT_GRANULARITY_KEY: _FUTURES_GRANULARITY,
            "clearing_house": "CME",
            "segment": "Rates",
            "notional": 40.0,
        },
        {
            _LIMIT_GRANULARITY_KEY: _FUTURES_GRANULARITY,
            "clearing_house": "ICE",
            "segment": "Credit",
            "notional": 60.0,
        },
    ]
    limits_cfg = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": "clearing_house",
                "entity_name": "CME",
                "limit_value": 0.35,
                "limit_kind": "percent_of_total",
            }
        ],
    }

    rows = _as_records(check_limits(exposures, limits_cfg))

    assert rows == [
        {
            "entity_type": "clearing_house",
            "entity_name": "cme",
            "limit_kind": "percent_of_total",
            "severity": "warning",
            "actual_value": pytest.approx(0.4),
            "limit_value": 0.35,
            "breach_amount": pytest.approx(0.05),
            "notes": None,
        }
    ]


def test_check_limits_is_deterministic_for_reversed_rows() -> None:
    exposures = [
        {"counterparty": "A", "notional": 10.0},
        {"counterparty": "B", "notional": 9.0},
        {"counterparty": "A", "notional": 8.0},
    ]
    limits_cfg = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": "counterparty",
                "entity_name": "A",
                "limit_value": 15.0,
                "limit_kind": "absolute_notional",
            }
        ],
    }

    first = _as_records(check_limits(exposures, limits_cfg))
    second = _as_records(check_limits(list(reversed(exposures)), limits_cfg))

    assert first == second


def test_check_limits_accepts_limits_config_object() -> None:
    exposures = [{"counterparty": "A", "notional": 11.0}]
    config = LimitsConfig.model_validate(
        {
            "schema_version": 1,
            "limits": [
                {
                    "entity_type": "counterparty",
                    "entity_name": "A",
                    "limit_value": 10.0,
                    "limit_kind": "absolute_notional",
                }
            ],
        }
    )

    rows = _as_records(check_limits(exposures, config))

    assert len(rows) == 1
    assert rows[0]["breach_amount"] == 1.0


def test_check_limits_validates_exposures_input() -> None:
    with pytest.raises(TypeError, match="exposures_df must be"):
        check_limits("not-a-table", {"schema_version": 1, "limits": []})

    with pytest.raises(ValueError, match="notional values must be numeric"):
        check_limits(
            [{"counterparty": "A", "notional": "oops"}],
            {
                "schema_version": 1,
                "limits": [
                    {
                        "entity_type": "counterparty",
                        "entity_name": "A",
                        "limit_value": 1.0,
                        "limit_kind": "absolute_notional",
                    }
                ],
            },
        )


def test_check_limits_validates_limits_cfg_input() -> None:
    with pytest.raises(TypeError, match="limits_cfg must be"):
        check_limits([{"counterparty": "A", "notional": 1.0}], 123)

    with pytest.raises(ValueError, match="limits_cfg is invalid"):
        check_limits(
            [{"counterparty": "A", "notional": 1.0}],
            {
                "schema_version": 1,
                "limits": [
                    {
                        "entity_type": "counterparty",
                        "entity_name": "A",
                        "limit_kind": "absolute_notional",
                    }
                ],
            },
        )


def test_write_limit_breaches_csv_writes_expected_rows(tmp_path: Path) -> None:
    exposures = [
        {"counterparty": "A", "notional": 11.0},
        {"counterparty": "B", "notional": 1.0},
    ]
    limits_cfg = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": "counterparty",
                "entity_name": "A",
                "limit_value": 10.0,
                "limit_kind": "absolute_notional",
            }
        ],
    }
    breaches = check_limits(exposures, limits_cfg)
    out = tmp_path / "limit_breaches.csv"

    write_limit_breaches_csv(breaches, out)

    assert out.exists()
    lines = out.read_text(encoding="utf-8").strip().splitlines()
    assert (
        lines[0]
        == "entity_type,entity_name,limit_kind,severity,actual_value,limit_value,breach_amount,notes"
    )
    assert "counterparty,a,absolute_notional,warning,11.0,10.0,1.0," in lines[1:]


def test_find_missing_limit_entities_returns_sorted_missing_targets() -> None:
    exposures = [
        {"counterparty": "A", "fcm": "FCM1", "notional": 11.0},
        {"counterparty": "B", "fcm": "FCM2", "notional": 1.0},
    ]
    limits_cfg = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": "counterparty",
                "entity_name": "A",
                "limit_value": 10.0,
                "limit_kind": "absolute_notional",
            },
            {
                "entity_type": "counterparty",
                "entity_name": "Missing Name",
                "limit_value": 10.0,
                "limit_kind": "absolute_notional",
            },
            {
                "entity_type": "fcm",
                "entity_name": "FCM Missing",
                "limit_value": 0.5,
                "limit_kind": "percent_of_total",
            },
        ],
    }

    missing = find_missing_limit_entities(exposures, limits_cfg)

    assert missing == [
        {"entity_type": "counterparty", "entity_name": "missing_name"},
        {"entity_type": "fcm", "entity_name": "fcm_missing"},
    ]


def test_check_limits_skips_disabled_limits_for_breaches_and_missing_entities() -> None:
    exposures = [{"counterparty": "A", "notional": 25.0}]
    limits_cfg = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": "counterparty",
                "entity_name": "A",
                "limit_value": 10.0,
                "limit_kind": "absolute_notional",
                "enabled": False,
            },
            {
                "entity_type": "counterparty",
                "entity_name": "Missing",
                "limit_value": 10.0,
                "limit_kind": "absolute_notional",
                "enabled": False,
            },
        ],
    }

    assert _as_records(check_limits(exposures, limits_cfg)) == []
    assert find_missing_limit_entities(exposures, limits_cfg) == []


def _counterparty_limit_exposure_rows(counterparty: str, notional: float) -> list[dict[str, Any]]:
    """Exercise the production reshaper with parsed counterparty totals."""

    return _build_limit_exposure_rows(
        {
            "all_programs": {
                "totals": [{"counterparty": counterparty, "Notional": notional}],
                "futures": [],
            }
        }
    )


def test_registered_alias_triggers_canonical_limit() -> None:
    limits_cfg = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": "counterparty",
                "entity_name": "bank_of_america",
                "limit_value": 100.0,
                "limit_kind": "absolute_notional",
                "severity": "fail",
            }
        ],
    }

    alias_exposures = _counterparty_limit_exposure_rows("Bank of America, NA", 200.0)
    alias_breaches = _as_records(check_limits(alias_exposures, limits_cfg))
    assert len(alias_breaches) == 1
    assert alias_breaches[0]["actual_value"] == 200.0
    assert alias_breaches[0]["breach_amount"] == 100.0
    assert alias_breaches[0]["severity"] == "fail"
    assert find_missing_limit_entities(alias_exposures, limits_cfg) == []

    canonical_exposures = _counterparty_limit_exposure_rows("Bank of America", 200.0)
    canonical_breaches = _as_records(check_limits(canonical_exposures, limits_cfg))
    assert canonical_breaches == alias_breaches

    unknown_exposures = _counterparty_limit_exposure_rows("Totally Unknown Counterparty", 200.0)
    assert _as_records(check_limits(unknown_exposures, limits_cfg)) == []
    assert find_missing_limit_entities(unknown_exposures, limits_cfg) == [
        {"entity_type": "counterparty", "entity_name": "bank_of_america"}
    ]


@pytest.mark.parametrize(
    "authoritative, expected_actual", [(200_000_000.0, None), (300_000_000.0, 300_000_000.0)]
)
@pytest.mark.parametrize("sign", [1, -1])
def test_all_programs_absolute_cap_ignores_overlapping_variants(
    authoritative: float, expected_actual: float | None, sign: int
) -> None:
    parsed = {
        "ex_trend": {
            "totals": [{"counterparty": "Citibank", "Notional": 150_000_000.0}],
            "futures": [],
        },
        "trend": {"totals": [], "futures": []},
        "all_programs": {
            "totals": [{"counterparty": "Citibank", "Notional": sign * authoritative}],
            "futures": [],
        },
    }
    limits = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": "counterparty",
                "entity_name": "citibank",
                "limit_kind": "absolute_notional",
                "limit_value": 250_000_000.0,
                "severity": "fail",
            }
        ],
    }
    breaches = _as_records(check_limits(_build_limit_exposure_rows(parsed), limits))
    if expected_actual is None:
        assert breaches == []
    else:
        assert len(breaches) == 1
        assert breaches[0]["actual_value"] == expected_actual
        assert breaches[0]["breach_amount"] == 50_000_000.0
        assert breaches[0]["severity"] == "fail"


@pytest.mark.parametrize(
    "subset_notional, subset_matches_target",
    [
        (None, False),
        (10.0, False),
        (900_000_000.0, False),
        (-900_000_000.0, False),
        (900_000_000.0, True),
        (-900_000_000.0, True),
    ],
)
@pytest.mark.parametrize(
    "entity_type, entity_name",
    [
        ("counterparty", "Citibank"),
        ("fcm", "FCM A"),
        ("clearing_house", "CME"),
        ("segment", "Rates"),
        ("custom_group", "Group A"),
    ],
)
def test_all_programs_percentage_universe_is_subset_invariant(
    subset_notional: float | None,
    subset_matches_target: bool,
    entity_type: str,
    entity_name: str,
) -> None:
    authoritative = {
        "totals": [
            {"counterparty": "Citibank", "Notional": 200.0},
            {"counterparty": "Other", "Notional": -800.0},
        ],
        "futures": [
            {
                "fcm": "FCM A",
                "clearing_house": "CME",
                "class": "Rates",
                "custom_group": "Group A",
                "notional": 40.0,
            },
            {
                "fcm": "FCM B",
                "clearing_house": "ICE",
                "class": "Credit",
                "custom_group": "Group B",
                "notional": -60.0,
            },
        ],
    }
    parsed = {}
    if subset_notional is not None:
        parsed.update(
            {
                name: {
                    "totals": [
                        {
                            "counterparty": "Citibank" if subset_matches_target else "Other",
                            "Notional": subset_notional,
                        }
                    ],
                    "futures": [
                        {
                            "fcm": "FCM A" if subset_matches_target else "FCM B",
                            "clearing_house": "CME" if subset_matches_target else "ICE",
                            "class": "Rates" if subset_matches_target else "Credit",
                            "custom_group": "Group A" if subset_matches_target else "Group B",
                            "notional": subset_notional,
                        }
                    ],
                }
                for name in ("ex_trend", "trend")
            }
        )
    # Select the authoritative key, even when reporting subsets appear first.
    parsed["all_programs"] = authoritative
    actual_share = 0.2 if entity_type == "counterparty" else 0.4
    limits = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": entity_type,
                "entity_name": entity_name,
                "limit_kind": "percent_of_total",
                "limit_value": actual_share - 0.05,
            }
        ],
    }
    rows = _build_limit_exposure_rows(parsed)
    breaches = _as_records(check_limits(rows, limits))
    assert len(breaches) == 1
    assert breaches[0]["actual_value"] == pytest.approx(actual_share)
    assert breaches[0]["breach_amount"] == pytest.approx(0.05)
    assert rows == _build_limit_exposure_rows({"all_programs": authoritative})
    limits["limits"][0]["limit_value"] = actual_share + 0.05
    assert _as_records(check_limits(rows, limits)) == []


def test_all_programs_missing_entities_cannot_be_satisfied_by_subset() -> None:
    parsed = {
        "all_programs": {"totals": [], "futures": []},
        "ex_trend": {
            "totals": [{"counterparty": "Citibank", "Notional": 150_000_000.0}],
            "futures": [
                {
                    "fcm": "FCM A",
                    "clearing_house": "CME",
                    "class": "Rates",
                    "custom_group": "Group A",
                    "notional": 40.0,
                }
            ],
        },
    }
    limits = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": kind,
                "entity_name": name,
                "limit_kind": "absolute_notional",
                "limit_value": 250_000_000.0,
            }
            for kind, name in [
                ("counterparty", "citibank"),
                ("fcm", "fcm_a"),
                ("clearing_house", "cme"),
                ("segment", "rates"),
                ("custom_group", "group_a"),
            ]
        ],
    }
    rows = _build_limit_exposure_rows(parsed)
    assert rows == []
    assert find_missing_limit_entities(rows, limits) == sorted(
        [
            {"entity_type": entry["entity_type"], "entity_name": entry["entity_name"]}
            for entry in limits["limits"]
        ],
        key=lambda row: row["entity_type"],
    )


@pytest.mark.parametrize("parsed", [{}, {"ex_trend": {"totals": [], "futures": []}}])
def test_all_programs_authority_is_required(parsed: dict[str, dict[str, Any]]) -> None:
    with pytest.raises(ValueError, match="Missing authoritative All Programs"):
        _build_limit_exposure_rows(parsed)


@pytest.mark.parametrize(
    ("entity_type", "canonical", "display_name", "alias"),
    [
        ("clearing_house", "cme", "CME", "CME Clearing House"),
        ("clearing_house", "ice", "ICE", "ICE Clear U.S."),
        ("fcm", "bank_of_america", "Bank of America", "Bank of America, NA"),
    ],
)
@pytest.mark.parametrize(
    ("limit_kind", "limit_value", "actual_value"),
    [("absolute_notional", 35.0, 60.0), ("percent_of_total", 0.35, 0.6)],
)
def test_registered_futures_alias_has_identical_limits_and_missing_results(
    entity_type: str,
    canonical: str,
    display_name: str,
    alias: str,
    limit_kind: str,
    limit_value: float,
    actual_value: float,
) -> None:
    policy = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": entity_type,
                "entity_name": canonical,
                "limit_value": limit_value,
                "limit_kind": limit_kind,
                "severity": "fail",
            }
        ],
    }
    other_row = {entity_type: "Unregistered Other Entity", "notional": 40.0}
    canonical_rows = [{entity_type: display_name, "notional": 60.0}, other_row]
    alias_rows = [{entity_type: alias, "notional": 60.0}, other_row]
    expected = _as_records(check_limits(canonical_rows, policy))
    assert expected == [
        {
            "entity_type": entity_type,
            "entity_name": canonical,
            "limit_kind": limit_kind,
            "severity": "fail",
            "actual_value": pytest.approx(actual_value),
            "limit_value": limit_value,
            "breach_amount": pytest.approx(actual_value - limit_value),
            "notes": None,
        }
    ]
    assert _as_records(check_limits(alias_rows, policy)) == expected
    assert (
        find_missing_limit_entities(alias_rows, policy)
        == find_missing_limit_entities(canonical_rows, policy)
        == []
    )


@pytest.mark.parametrize("entity_type", ["fcm", "clearing_house"])
def test_new_registry_futures_aliases_aggregate_into_one_limit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, entity_type: str
) -> None:
    registry_path = tmp_path / "config" / "name_registry.yml"
    registry_path.parent.mkdir()
    registry_path.write_text(
        "schema_version: 1\nentries:\n"
        "  - canonical_key: registered_futures_entity\n"
        "    display_name: Registered Futures Entity\n"
        "    aliases:\n      - First Futures Alias\n      - Second Futures Alias\n",
        encoding="utf-8",
    )
    monkeypatch.chdir(tmp_path)
    policy = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": entity_type,
                "entity_name": "registered_futures_entity",
                "limit_value": 0.35,
                "limit_kind": "percent_of_total",
            }
        ],
    }
    other_row = {entity_type: "Unregistered Other Entity", "notional": 40.0}
    canonical_rows = [
        {entity_type: "registered_futures_entity", "notional": 60.0},
        other_row,
    ]
    expected = _as_records(check_limits(canonical_rows, policy))
    assert len(expected) == 1
    assert expected[0]["actual_value"] == pytest.approx(0.6)
    assert expected[0]["breach_amount"] == pytest.approx(0.25)
    for rows in (
        [{entity_type: "Registered Futures Entity", "notional": 60.0}, other_row],
        [
            {entity_type: "First Futures Alias", "notional": 30.0},
            {entity_type: "Second Futures Alias", "notional": 30.0},
            other_row,
        ],
    ):
        assert _as_records(check_limits(rows, policy)) == expected
        assert find_missing_limit_entities(rows, policy) == []


@pytest.mark.parametrize(
    "clearing_rows",
    [
        [("CME", 60.0), ("ICE", 40.0)],
        [("CME Clearing House", 60.0), ("ICE Clear U.S.", 40.0)],
        [("CME", 30.0), ("CME Clearing House", 30.0), ("ICE Clear U.S.", 40.0)],
    ],
    ids=["canonical", "alias", "split-alias"],
)
def test_shipped_cme_limit_breaches_for_clearing_house_aliases(
    clearing_rows: list[tuple[str, float]],
) -> None:
    policy = load_limits_config()
    rows = _build_limit_exposure_rows(
        {
            "all_programs": {
                "totals": [{"counterparty": "Citibank", "Notional": 1000.0}],
                "futures": [
                    {"clearing_house": label, "notional": amount} for label, amount in clearing_rows
                ],
            }
        }
    )

    breaches = _as_records(check_limits(rows, policy))

    assert len(breaches) == 1
    assert breaches[0]["entity_type"] == "clearing_house"
    assert breaches[0]["entity_name"] == "cme"
    assert breaches[0]["limit_kind"] == "percent_of_total"
    assert breaches[0]["limit_value"] == 0.35
    assert breaches[0]["actual_value"] == pytest.approx(0.6)
    assert breaches[0]["breach_amount"] == pytest.approx(0.25)
    assert find_missing_limit_entities(rows, policy) == []


def test_clearing_house_split_alias_aggregation_preserves_percent_breach() -> None:
    policy = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": "clearing_house",
                "entity_name": "cme",
                "limit_value": 0.35,
                "limit_kind": "percent_of_total",
            }
        ],
    }
    canonical = [
        {"clearing_house": "CME", "notional": 60.0},
        {"clearing_house": "ICE", "notional": 40.0},
    ]
    split = [
        {"clearing_house": "CME", "notional": 30.0},
        {"clearing_house": "CME Clearing House", "notional": 30.0},
        {"clearing_house": "ICE Clear U.S.", "notional": 40.0},
    ]
    expected = _as_records(check_limits(canonical, policy))
    assert expected[0]["actual_value"] == pytest.approx(0.6)
    assert _as_records(check_limits(split, policy)) == expected
    assert find_missing_limit_entities(split, policy) == []


@pytest.mark.parametrize("disabled_alias_present", [False, True])
def test_unknown_clearing_house_alias_and_disabled_limit_preserve_missing_policy(
    disabled_alias_present: bool,
) -> None:
    policy = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": "clearing_house",
                "entity_name": "cme",
                "limit_value": 35.0,
                "limit_kind": "absolute_notional",
            },
            {
                "entity_type": "clearing_house",
                "entity_name": "ice",
                "limit_value": 1.0,
                "limit_kind": "absolute_notional",
                "enabled": False,
            },
        ],
    }
    rows = [{"clearing_house": "Unregistered Clearing House", "notional": 60.0}]
    if disabled_alias_present:
        rows.append({"clearing_house": "ICE Clear U.S.", "notional": 40.0})
    assert _as_records(check_limits(rows, policy)) == []
    assert find_missing_limit_entities(rows, policy) == [
        {"entity_type": "clearing_house", "entity_name": "cme"}
    ]


@pytest.mark.parametrize("entity_type", ["segment", "custom_group"])
def test_policy_group_aliases_are_not_resolved_through_registry(entity_type: str) -> None:
    policy = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": entity_type,
                "entity_name": "cme",
                "limit_value": 35.0,
                "limit_kind": "absolute_notional",
            }
        ],
    }
    rows = [{entity_type: "CME Clearing House", "notional": 60.0}]
    assert _as_records(check_limits(rows, policy)) == []
    assert find_missing_limit_entities(rows, policy) == [
        {"entity_type": entity_type, "entity_name": "cme"}
    ]


@pytest.fixture
def boundary_parsed_workbooks(
    tmp_path: Path, registered_alias_workbook: Path
) -> dict[str, dict[str, Any]]:
    """Exercise actual workbook parsing before the two numerical consumers."""
    from openpyxl import load_workbook

    parsed = {}
    for variant, citi, alias, other, futures_scale in (
        ("all_programs", 60.0, -40.0, 40.0, 1.0),
        ("ex_trend", 600.0, -300.0, 900.0, 10.0),
        ("trend", 6000.0, -3000.0, 9000.0, 100.0),
    ):
        path = tmp_path / f"boundary-{variant.replace('_', '-')}.xlsx"
        workbook = load_workbook(registered_alias_workbook)
        sheet = workbook.active
        assert sheet is not None
        for row, name, equity, treasury in (
            (12, "Citibank", citi, 0.0),
            (13, "Citigroup", alias, 0.0),
            (14, "Other", 0.0, other),
        ):
            sheet.cell(row, 4, name)
            sheet.cell(row, 7, treasury)
            sheet.cell(row, 8, equity)
            sheet.cell(row, 12, equity + treasury)
        for column, header in enumerate(
            ("Account", "Description", "Class", "FCM", "Clearing House", "Notional"), 4
        ):
            sheet.cell(24, column, header)
        for row, fcm, house, amount in (
            (25, "Bank of America, NA", "CME Clearing House", 60.0),
            (26, "Unknown FCM", "ICE Clear U.S.", -40.0),
        ):
            for column, value in enumerate(
                (f"acct-{row}", "Future", "Equity", fcm, house, amount * futures_scale), 4
            ):
                sheet.cell(row, column, value)
        sheet.cell(27, 4, "risk exclusive of the trend positions")
        workbook.save(path)
        workbook.close()
        parsed[variant] = {
            "totals": parse_fcm_totals(path),
            "futures": parse_futures_detail(path),
        }
    return parsed


@pytest.mark.parametrize(
    "entity_type, canonical", [("fcm", "bank_of_america"), ("clearing_house", "cme")]
)
def test_boundary_audit_alias_only_parser_to_limit(
    boundary_parsed_workbooks: dict[str, dict[str, Any]], entity_type: str, canonical: str
) -> None:
    policy = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": entity_type,
                "entity_name": canonical,
                "limit_kind": "percent_of_total",
                "limit_value": 0.5,
            }
        ],
    }
    rows = _build_limit_exposure_rows(boundary_parsed_workbooks)
    breaches = _as_records(check_limits(rows, policy))
    assert len(breaches) == 1, "parsed registered alias must reach the canonical cap"
    assert breaches[0]["entity_name"] == canonical
    assert breaches[0]["actual_value"] == pytest.approx(0.6)
    assert breaches[0]["breach_amount"] == pytest.approx(0.1)
    assert find_missing_limit_entities(rows, policy) == []


def test_boundary_audit_overlap_parser_to_limit(
    boundary_parsed_workbooks: dict[str, dict[str, Any]],
) -> None:
    policy = {
        "schema_version": 1,
        "limits": [
            {
                "entity_type": "counterparty",
                "entity_name": "citibank",
                "limit_kind": "absolute_notional",
                "limit_value": 125.0,
            },
            {
                "entity_type": "counterparty",
                "entity_name": "citibank",
                "limit_kind": "percent_of_total",
                "limit_value": 0.5,
            },
            {
                "entity_type": "counterparty",
                "entity_name": "subset_only",
                "limit_kind": "absolute_notional",
                "limit_value": 1.0,
            },
        ],
    }
    # A subset-only entity must supply neither denominator nor coverage.
    boundary_parsed_workbooks["ex_trend"]["totals"] = [
        {"counterparty": "Citigroup", "Notional": 900.0},
        {"counterparty": "Subset Only", "Notional": 9000.0},
    ]
    for parsed in (
        boundary_parsed_workbooks,
        {"all_programs": boundary_parsed_workbooks["all_programs"]},
    ):
        rows = _build_limit_exposure_rows(parsed)
        breaches = _as_records(check_limits(rows, policy))
        assert len(breaches) == 1, "overlapping subsets must not manufacture an absolute breach"
        assert breaches[0]["limit_kind"] == "percent_of_total"
        assert breaches[0]["actual_value"] == pytest.approx(5 / 7)
        assert breaches[0]["breach_amount"] == pytest.approx(5 / 7 - 0.5)
        assert find_missing_limit_entities(rows, policy) == [
            {"entity_type": "counterparty", "entity_name": "subset_only"}
        ]
    policy["limits"][0]["limit_value"] = 75.0
    breaches = _as_records(check_limits(rows, policy))
    assert breaches[0]["actual_value"] == 100.0
    assert breaches[0]["breach_amount"] == 25.0


def test_boundary_audit_partitioned_concentration_control(
    tmp_path: Path, boundary_parsed_workbooks: dict[str, dict[str, Any]]
) -> None:
    metrics = _compute_and_write_concentration_metrics(
        parsed_by_variant=boundary_parsed_workbooks, run_dir=tmp_path
    )
    by_partition = {(row["variant"], row["segment"]): row for row in metrics}
    assert set(by_partition) == {
        (variant, segment)
        for variant in ("all_programs", "ex_trend")
        for segment in ("TIPS", "Treasury", "Equity", "Commodity", "Currency", "total")
    }, "concentration must retain variant and segment; Trend has no parsed totals"
    for variant, expected_total in (("all_programs", 29 / 49), ("ex_trend", 0.5)):
        for segment, expected_hhi in (
            ("total", expected_total),
            ("Equity", 1.0),
            ("Treasury", 1.0),
            ("TIPS", 0.0),
            ("Commodity", 0.0),
            ("Currency", 0.0),
        ):
            row = by_partition[variant, segment]
            assert row["hhi"] == pytest.approx(expected_hhi)
            for metric in ("top5_share", "top10_share"):
                assert row[metric] == pytest.approx(1.0 if expected_hhi else 0.0)
