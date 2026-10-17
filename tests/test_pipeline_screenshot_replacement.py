from __future__ import annotations

import os
import sys
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Literal
from zipfile import ZipFile

import pytest

import counter_risk.pipeline.run as run_module
from counter_risk.config import WorkflowConfig
from counter_risk.pipeline.run import PptProcessingResult, PptProcessingStatus


@dataclass
class _ScreenshotBackendCall:
    source: Path
    output: Path
    mapping: dict[str, Path]


def _write_placeholder(path: Path, *, payload: bytes = b"fixture") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)


def _write_pptx_with_external_link(path: Path, target: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    rels_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" Target="{target}" TargetMode="External"/>
</Relationships>
"""
    with ZipFile(path, "w") as archive:
        archive.writestr("ppt/slides/_rels/slide1.xml.rels", rels_xml)


def _read_relationship_targets(path: Path, rel_path: str) -> list[str]:
    with ZipFile(path) as archive:
        xml_bytes = archive.read(rel_path)
    root = ET.fromstring(xml_bytes)
    return [
        rel.attrib.get("Target", "").strip()
        for rel in root.findall(
            ".//{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"
        )
        if rel.attrib.get("Target", "").strip()
    ]


def _build_config(
    *,
    tmp_path: Path,
    enable_screenshot_replacement: bool,
    screenshot_replacement_implementation: Literal["zip", "python-pptx"] = "zip",
    screenshot_inputs: dict[str, Path] | None = None,
) -> WorkflowConfig:
    inputs_dir = tmp_path / "inputs"
    files = {
        "all_programs": inputs_dir / "all_programs.xlsx",
        "ex_trend": inputs_dir / "ex_trend.xlsx",
        "trend": inputs_dir / "trend.xlsx",
        "hist_all": inputs_dir / "hist_all.xlsx",
        "hist_ex": inputs_dir / "hist_ex.xlsx",
        "hist_llc": inputs_dir / "hist_llc.xlsx",
        "monthly_pptx": inputs_dir / "monthly.pptx",
    }
    for path in files.values():
        _write_placeholder(path)

    return WorkflowConfig(
        mosers_all_programs_xlsx=files["all_programs"],
        mosers_ex_trend_xlsx=files["ex_trend"],
        mosers_trend_xlsx=files["trend"],
        hist_all_programs_3yr_xlsx=files["hist_all"],
        hist_ex_llc_3yr_xlsx=files["hist_ex"],
        hist_llc_3yr_xlsx=files["hist_llc"],
        monthly_pptx=files["monthly_pptx"],
        output_root=tmp_path / "ignored-output-root",
        enable_screenshot_replacement=enable_screenshot_replacement,
        screenshot_replacement_implementation=screenshot_replacement_implementation,
        screenshot_inputs=screenshot_inputs or {},
    )


def test_write_outputs_calls_zip_backend_once_when_enabled(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir(parents=True)
    image_1 = tmp_path / "screenshots" / "slide_1.png"
    image_2 = tmp_path / "screenshots" / "slide_2.png"
    _write_placeholder(image_1, payload=b"img-1")
    _write_placeholder(image_2, payload=b"img-2")

    config = _build_config(
        tmp_path=tmp_path,
        enable_screenshot_replacement=True,
        screenshot_replacement_implementation="zip",
        screenshot_inputs={"slide2": image_2, "slide1": image_1},
    )

    calls: list[_ScreenshotBackendCall] = []

    def _fake_zip_backend(source: Path, output: Path, mapping: dict[str, Path]) -> None:
        calls.append(_ScreenshotBackendCall(source=source, output=output, mapping=mapping))
        output.write_bytes(source.read_bytes() + b"-replaced")

    monkeypatch.setattr(run_module, "_replace_screenshots_with_zip_backend", _fake_zip_backend)
    monkeypatch.setattr(
        run_module,
        "_refresh_ppt_links",
        lambda _path: PptProcessingResult(status=PptProcessingStatus.SUCCESS),
    )

    warnings: list[str] = []
    output_paths, _ = run_module._write_outputs(
        run_dir=run_dir,
        config=config,
        as_of_date=date(2026, 2, 13),
        warnings=warnings,
    )

    assert len(calls) == 1
    assert calls[0].output == (
        run_dir / "Monthly Counterparty Exposure Report (Master) - 2026-02-13.pptx"
    )
    assert list(calls[0].mapping) == ["slide1", "slide2"]
    assert len(calls[0].mapping) == 2
    assert all(path.suffix.lower() == ".png" for path in calls[0].mapping.values())
    master = run_dir / "Monthly Counterparty Exposure Report (Master) - 2026-02-13.pptx"
    distribution = run_dir / "Monthly Counterparty Exposure Report - 2026-02-13.pptx"
    assert master in output_paths
    assert distribution in output_paths
    assert output_paths.index(master) < output_paths.index(distribution)
    assert master.exists()
    assert distribution.exists()
    assert all("not implemented" not in warning for warning in warnings)
    assert all("disabled; copied source deck unchanged" not in warning for warning in warnings)


def test_write_outputs_routes_python_pptx_backend_when_selected(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir(parents=True)
    image_1 = tmp_path / "screenshots" / "all_programs.png"
    _write_placeholder(image_1, payload=b"img-1")

    config = _build_config(
        tmp_path=tmp_path,
        enable_screenshot_replacement=True,
        screenshot_replacement_implementation="python-pptx",
        screenshot_inputs={"All Programs": image_1},
    )

    called = {"python_pptx": 0, "zip": 0}

    def _fake_python_pptx_backend(source: Path, output: Path, mapping: dict[str, Path]) -> None:
        _ = mapping
        called["python_pptx"] += 1
        output.write_bytes(source.read_bytes() + b"-replaced")

    def _fake_zip_backend(_source: Path, _output: Path, _mapping: dict[str, Path]) -> None:
        called["zip"] += 1

    monkeypatch.setattr(
        run_module,
        "_replace_screenshots_with_python_pptx_backend",
        _fake_python_pptx_backend,
    )
    monkeypatch.setattr(run_module, "_replace_screenshots_with_zip_backend", _fake_zip_backend)
    monkeypatch.setattr(
        run_module,
        "_refresh_ppt_links",
        lambda _path: PptProcessingResult(status=PptProcessingStatus.SUCCESS),
    )

    warnings: list[str] = []
    run_module._write_outputs(
        run_dir=run_dir,
        config=config,
        as_of_date=date(2026, 2, 13),
        warnings=warnings,
    )

    assert called["python_pptx"] == 1
    assert called["zip"] == 0
    assert all("not implemented" not in warning for warning in warnings)


def test_write_outputs_derives_distribution_from_refreshed_master(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir(parents=True)

    config = _build_config(
        tmp_path=tmp_path,
        enable_screenshot_replacement=False,
    )
    config.monthly_pptx.write_bytes(b"base-ppt")

    def _refresh(path: Path) -> PptProcessingResult:
        path.write_bytes(path.read_bytes() + b"-refreshed")
        return PptProcessingResult(status=PptProcessingStatus.SUCCESS)

    monkeypatch.setattr(run_module, "_refresh_ppt_links", _refresh)

    output_paths, _ = run_module._write_outputs(
        run_dir=run_dir,
        config=config,
        as_of_date=date(2026, 2, 13),
        warnings=[],
    )

    master = run_dir / "Monthly Counterparty Exposure Report (Master) - 2026-02-13.pptx"
    distribution = run_dir / "Monthly Counterparty Exposure Report - 2026-02-13.pptx"
    assert master in output_paths
    assert distribution in output_paths
    assert output_paths.index(master) < output_paths.index(distribution)
    assert master.read_bytes().endswith(b"-refreshed")
    assert distribution.read_bytes() == master.read_bytes()


def test_write_outputs_raises_when_master_generation_drops_external_links(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir(parents=True)
    image_1 = tmp_path / "screenshots" / "slide_1.png"
    _write_placeholder(image_1, payload=b"img-1")

    config = _build_config(
        tmp_path=tmp_path,
        enable_screenshot_replacement=True,
        screenshot_replacement_implementation="zip",
        screenshot_inputs={"slide1": image_1},
    )
    _write_pptx_with_external_link(config.monthly_pptx, "X:\\linked\\book.xlsx")

    def _drop_links(_source: Path, output: Path, _mapping: dict[str, Path]) -> None:
        with ZipFile(output, "w") as archive:
            archive.writestr("ppt/slides/slide1.xml", "<p:sld/>")

    monkeypatch.setattr(run_module, "_replace_screenshots_with_zip_backend", _drop_links)
    monkeypatch.setattr(
        run_module,
        "_refresh_ppt_links",
        lambda _path: PptProcessingResult(status=PptProcessingStatus.SUCCESS),
    )

    with pytest.raises(RuntimeError, match="did not preserve linked chart references"):
        run_module._write_outputs(
            run_dir=run_dir,
            config=config,
            as_of_date=date(2026, 2, 13),
            warnings=[],
        )


def test_write_outputs_keeps_master_external_links_when_backend_preserves_relationships(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir(parents=True)
    image_1 = tmp_path / "screenshots" / "slide_1.png"
    _write_placeholder(image_1, payload=b"img-1")

    config = _build_config(
        tmp_path=tmp_path,
        enable_screenshot_replacement=True,
        screenshot_replacement_implementation="zip",
        screenshot_inputs={"slide1": image_1},
    )
    _write_pptx_with_external_link(config.monthly_pptx, "X:\\linked\\book.xlsx")

    def _preserve_links(source: Path, output: Path, _mapping: dict[str, Path]) -> None:
        output.write_bytes(source.read_bytes())

    monkeypatch.setattr(run_module, "_replace_screenshots_with_zip_backend", _preserve_links)
    monkeypatch.setattr(
        run_module,
        "_refresh_ppt_links",
        lambda _path: PptProcessingResult(status=PptProcessingStatus.SUCCESS),
    )

    output_paths, _ = run_module._write_outputs(
        run_dir=run_dir,
        config=config,
        as_of_date=date(2026, 2, 13),
        warnings=[],
    )

    master = run_dir / "Monthly Counterparty Exposure Report (Master) - 2026-02-13.pptx"
    distribution = run_dir / "Monthly Counterparty Exposure Report - 2026-02-13.pptx"
    assert master in output_paths
    assert distribution in output_paths
    assert output_paths.index(master) < output_paths.index(distribution)


def test_derive_distribution_ppt_removes_external_link_relationships_from_master(
    tmp_path: Path,
) -> None:
    master = tmp_path / "master.pptx"
    distribution = tmp_path / "distribution.pptx"
    with ZipFile(master, "w") as archive:
        archive.writestr(
            "ppt/slides/_rels/slide1.xml.rels",
            """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" Target="X:\\linked\\book.xlsx" TargetMode="External"/>
</Relationships>
""",
        )
        archive.writestr("ppt/media/image1.png", b"png-data")

    run_module._derive_distribution_ppt(
        master_pptx_path=master, distribution_pptx_path=distribution
    )

    assert run_module._list_external_link_targets(master) == {"X:\\linked\\book.xlsx"}
    assert run_module._list_external_link_targets(distribution) == set()
    assert (tmp_path / "distribution.pptx").exists()
    assert _read_relationship_targets(distribution, "ppt/slides/_rels/slide1.xml.rels") == [
        "../media/image1.png"
    ]
    with ZipFile(distribution) as archive:
        assert archive.read("ppt/media/image1.png") == b"png-data"


def test_resolve_screenshot_input_mapping_rejects_non_png_paths(tmp_path: Path) -> None:
    text_path = tmp_path / "screenshots" / "slide_1.jpg"
    _write_placeholder(text_path, payload=b"not-png")
    config = _build_config(
        tmp_path=tmp_path,
        enable_screenshot_replacement=True,
        screenshot_replacement_implementation="zip",
        screenshot_inputs={"slide1": text_path},
    )

    with pytest.raises(ValueError, match="must point to a PNG file"):
        run_module._resolve_screenshot_input_mapping(config)


def test_resolve_screenshot_input_mapping_sorts_by_normalized_keys(tmp_path: Path) -> None:
    image_1 = tmp_path / "screenshots" / "slide_1.png"
    image_2 = tmp_path / "screenshots" / "slide_2.png"
    _write_placeholder(image_1, payload=b"img-1")
    _write_placeholder(image_2, payload=b"img-2")
    config = _build_config(
        tmp_path=tmp_path,
        enable_screenshot_replacement=False,
        screenshot_inputs={"  slide2  ": image_2, "slide1": image_1},
    )

    mapping = run_module._resolve_screenshot_input_mapping(config)

    assert list(mapping.keys()) == ["slide1", "slide2"]


def test_resolve_screenshot_input_mapping_rejects_duplicate_normalized_keys(tmp_path: Path) -> None:
    image_1 = tmp_path / "screenshots" / "slide_1.png"
    image_2 = tmp_path / "screenshots" / "slide_2.png"
    _write_placeholder(image_1, payload=b"img-1")
    _write_placeholder(image_2, payload=b"img-2")
    config = _build_config(
        tmp_path=tmp_path,
        enable_screenshot_replacement=True,
        screenshot_inputs={"slide1": image_1, " slide1 ": image_2},
    )

    with pytest.raises(ValueError, match="duplicated after normalization"):
        run_module._resolve_screenshot_input_mapping(config)


def test_resolve_screenshot_input_mapping_rejects_apostrophe_variant_duplicates(
    tmp_path: Path,
) -> None:
    image_1 = tmp_path / "screenshots" / "slide_1.png"
    image_2 = tmp_path / "screenshots" / "slide_2.png"
    _write_placeholder(image_1, payload=b"img-1")
    _write_placeholder(image_2, payload=b"img-2")
    # ASCII apostrophe U+0027 and curly right apostrophe U+2019 are the same canonical key
    ascii_apos_key = "Goldman Sachs Int'l"
    curly_apos_key = "Goldman Sachs Int’l"
    config = _build_config(
        tmp_path=tmp_path,
        enable_screenshot_replacement=True,
        screenshot_inputs={ascii_apos_key: image_1, curly_apos_key: image_2},
    )

    with pytest.raises(ValueError, match="duplicated after normalization"):
        run_module._resolve_screenshot_input_mapping(config)


def test_resolve_screenshot_input_mapping_rejects_dash_variant_duplicates(
    tmp_path: Path,
) -> None:
    image_1 = tmp_path / "screenshots" / "slide_1.png"
    image_2 = tmp_path / "screenshots" / "slide_2.png"
    _write_placeholder(image_1, payload=b"img-1")
    _write_placeholder(image_2, payload=b"img-2")
    config = _build_config(
        tmp_path=tmp_path,
        enable_screenshot_replacement=True,
        # ASCII hyphen vs en-dash (U+2013) → same canonical key
        screenshot_inputs={"Ex-Trend": image_1, "Ex–Trend": image_2},
    )

    with pytest.raises(ValueError, match="duplicated after normalization"):
        run_module._resolve_screenshot_input_mapping(config)


def test_resolve_screenshot_input_mapping_canonicalizes_key_stored_in_mapping(
    tmp_path: Path,
) -> None:
    image = tmp_path / "screenshots" / "slide_1.png"
    _write_placeholder(image, payload=b"img-1")
    # Curly apostrophe U+2019 and extra spaces get normalized to canonical form
    raw_key = "Goldman  Sachs  Int’l"
    config = _build_config(
        tmp_path=tmp_path,
        enable_screenshot_replacement=False,
        screenshot_inputs={raw_key: image},
    )

    mapping = run_module._resolve_screenshot_input_mapping(config)

    # Canonical form: ASCII apostrophe, collapsed whitespace
    assert list(mapping.keys()) == ["Goldman Sachs Int'l"]


@pytest.mark.parametrize(
    "native_office",
    [
        pytest.param(False, id="portable"),
        pytest.param(
            True,
            id="office",
            marks=[
                pytest.mark.office,
                pytest.mark.skipif(
                    sys.platform != "win32" or os.environ.get("COUNTER_RISK_OFFICE_SMOKE") != "1",
                    reason="Requires Windows Excel/PowerPoint COM and COUNTER_RISK_OFFICE_SMOKE=1",
                ),
            ],
        ),
    ],
)
def test_office_cash_change_updates_final_ppt_media_preserving_inputs(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    cprs_workflow_config: WorkflowConfig,
    native_office: bool,
) -> None:
    """Check both delivered PPTs with native Office or a cash-reading renderer."""
    import hashlib
    import json
    from io import BytesIO

    import yaml
    from PIL import Image, ImageChops
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if not native_office:
        from openpyxl import load_workbook
        from PIL import ImageDraw

        from counter_risk.integrations import excel_com

        def export(*, workbook_path: Path, sheet_name: str, output_png: Path) -> Path:
            # Read the actual exporter-bound cells so the original-workbook
            # route produces stale pixels and fails the shared smoke assertions.
            workbook = load_workbook(workbook_path, data_only=True)
            try:
                cash = [row[3] for row in workbook[sheet_name].values if row[2] == "CIBC"]
            finally:
                workbook.close()
            output_png.parent.mkdir(parents=True, exist_ok=True)
            image = Image.new("RGB", (320, 80), "white")
            ImageDraw.Draw(image).text((10, 10), f"{sheet_name}: CIBC cash {cash}", fill="black")
            image.save(output_png)
            return output_png

        monkeypatch.setattr(excel_com, "export_worksheet_range_as_png", export)
        monkeypatch.setattr(
            run_module,
            "_refresh_ppt_links",
            lambda _: PptProcessingResult(status=PptProcessingStatus.SUCCESS),
        )

    config = cprs_workflow_config
    original_hashes = {
        path: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in (tmp_path / "inputs").glob("*.xlsx")
    }
    media_by_run: list[dict[tuple[str, int], bytes]] = []
    assert config.cash_source_path is not None
    for index, amount in enumerate((1234567, 7654321)):
        config.cash_source_path.write_text(f"counterparty,cash_value\nCIBC,{amount}\n")
        config_path = tmp_path / f"office-{index}.yml"
        config_path.write_text(yaml.safe_dump(config.model_dump(mode="json")))
        run_dir = run_module.run_pipeline(config_path, output_dir=tmp_path / f"office-{index}")
        media: dict[tuple[str, int], bytes] = {}
        for kind, filename in (
            ("master", "Monthly Counterparty Exposure Report (Master) - 2025-12-31.pptx"),
            ("distribution", "Monthly Counterparty Exposure Report - 2025-12-31.pptx"),
        ):
            deck = run_dir / filename
            assert deck.exists(), f"Office did not deliver {kind}: inspect run warnings"
            presentation = Presentation(deck)
            for slide_number in (1, 2):
                pictures = [
                    shape
                    for shape in presentation.slides[slide_number - 1].shapes
                    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                ]
                assert pictures, f"Missing CPRS image on {kind} slide {slide_number}"
                media[kind, slide_number] = pictures[0].image.blob
                assert media[kind, slide_number].startswith(b"\x89PNG\r\n\x1a\n")
                table = "ch" if slide_number == 1 else "fcm"
                generated = run_dir / "_screenshots" / f"all_programs_{table}.png"
                assert (
                    media[kind, slide_number] == generated.read_bytes()
                ), f"Final {kind} slide {slide_number} must embed the generated CPRS image"
                print(
                    index, kind, slide_number, hashlib.sha256(media[kind, slide_number]).hexdigest()
                )
        media_by_run.append(media)
        assert {
            path: hashlib.sha256(path.read_bytes()).hexdigest() for path in original_hashes
        } == original_hashes
    assert media_by_run[0].keys() == media_by_run[1].keys()
    for key in media_by_run[0]:
        # PNG metadata or compression changes alone do not prove a cash change
        # reached the visible table. Compare decoded pixels.
        with (
            Image.open(BytesIO(media_by_run[0][key])) as before,
            Image.open(BytesIO(media_by_run[1][key])) as after,
        ):
            assert (
                before.size != after.size
                or ImageChops.difference(before.convert("RGB"), after.convert("RGB")).getbbox()
                is not None
            ), f"Stale CPRS media: {key}"
    print(
        "unchanged input SHA-256", {str(path): digest for path, digest in original_hashes.items()}
    )
    (tmp_path / "cash-media-evidence.json").write_text(
        json.dumps(
            {
                "renderer": "native-office" if native_office else "portable-test-double",
                "cash_amounts": [1234567, 7654321],
                "unchanged_input_sha256": {
                    str(path.relative_to(tmp_path)): digest
                    for path, digest in original_hashes.items()
                },
                "embedded_png_sha256_by_run": [
                    {
                        f"{kind}/slide{slide}": hashlib.sha256(blob).hexdigest()
                        for (kind, slide), blob in media.items()
                    }
                    for media in media_by_run
                ],
                "generated_images_embedded": True,
                "decoded_pixels_changed": True,
            },
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
