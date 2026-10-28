"""Global pytest collection hooks for CI-friendly test selection."""

from __future__ import annotations

from pathlib import Path

import pytest


def pytest_addoption(parser: pytest.Parser) -> None:
    parser.addoption(
        "--slow-shard",
        default=None,
        metavar="I/N",
        help="run only shard I of N (1-based) of the tests selected by -m/paths, split round-robin "
        "over the sorted node IDs so the N shards are disjoint and together select every test",
    )


def parse_shard(spec: str) -> tuple[int, int]:
    """`I/N` -> (I, N); a malformed or out-of-range spec raises instead of selecting nothing."""
    try:
        index_text, count_text = spec.split("/")
        index, count = int(index_text), int(count_text)
    except ValueError as exc:
        raise pytest.UsageError(f"--slow-shard expects I/N, got {spec!r}") from exc
    if not 1 <= index <= count:
        raise pytest.UsageError(f"--slow-shard {spec!r}: need 1 <= I <= N")
    return index, count


class _SlowShardPlugin:
    """Keep shard I of N of what -m/paths already selected.

    A separate plugin so its hook can be `trylast`: marker deselection must have run first, or the
    shards would be cut from the whole suite instead of from the slow tests they are meant to split.
    """

    def __init__(self, index: int, count: int) -> None:
        self.index, self.count = index, count

    @pytest.hookimpl(trylast=True)
    def pytest_collection_modifyitems(
        self, config: pytest.Config, items: list[pytest.Item]
    ) -> None:
        ordered = sorted(items, key=lambda item: item.nodeid)
        keep = {
            id(item)
            for position, item in enumerate(ordered)
            if position % self.count == self.index - 1
        }
        dropped = [item for item in items if id(item) not in keep]
        items[:] = [item for item in items if id(item) in keep]
        config.hook.pytest_deselected(items=dropped)


def pytest_configure(config: pytest.Config) -> None:
    spec = config.getoption("--slow-shard")
    if spec is not None:
        index, count = parse_shard(spec)
        config.pluginmanager.register(_SlowShardPlugin(index, count), "slow-shard")


@pytest.fixture
def registered_alias_workbook(tmp_path: Path) -> Path:
    """Real CPRS-FCM input whose two labels represent one registered entity."""
    from openpyxl import Workbook

    path = tmp_path / "registered-aliases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "CPRS - FCM"
    for row, column, value in (
        (4, 4, "Position"),
        (5, 4, "Counterparty/ FCM"),
        (5, 7, "Nominal"),
        (5, 12, "12/31 Total"),
        (5, 13, "Notional change"),
        (6, 6, "TIPS"),
        (6, 7, "Treasury"),
        (6, 8, "Equity"),
        (6, 9, "Commodity"),
        (6, 10, "Currency"),
        (6, 12, "Notional"),
        (6, 13, "from prior month"),
        (10, 4, "Total by Counterparty/ FCM (This is not the legal obligation exposure)"),
        (23, 4, "FUTURES DETAIL"),
    ):
        sheet.cell(row, column, value)
    for row, name, notional in ((12, "Citibank", 60.0), (13, "Citigroup", -40.0)):
        sheet.cell(row, 4, name)
        sheet.cell(row, 8, notional)
        sheet.cell(row, 12, notional)
    workbook.save(path)
    workbook.close()
    return path


def pytest_collection_modifyitems(items: list[pytest.Item]) -> None:
    """Mark integration directory tests as slow for PR-gate runs."""
    slow = pytest.mark.slow
    for item in items:
        item_path = getattr(item, "path", None)
        if item_path is None:
            continue
        try:
            relative_parts = Path(item_path).parts
        except TypeError:
            continue
        if "integration" in relative_parts and "tests" in relative_parts:
            item.add_marker(slow)


@pytest.fixture
def cprs_workflow_config(tmp_path: Path):
    """Real reporting inputs copied away from the immutable fixture archive."""
    import shutil
    from datetime import date

    from counter_risk.config import WorkflowConfig

    fixtures = Path(__file__).parent / "fixtures"
    names = {
        "mosers_all_programs_xlsx": "MOSERS Counterparty Risk Summary 12-31-2025 - All Programs.xlsx",
        "mosers_ex_trend_xlsx": "MOSERS Counterparty Risk Summary 12-31-2025 - Ex Trend.xlsx",
        "mosers_trend_xlsx": "MOSERS Counterparty Risk Summary 12-31-2025 - Trend.xlsx",
        "hist_all_programs_3yr_xlsx": "Historical Counterparty Risk Graphs - All Programs 3 Year.xlsx",
        "hist_ex_llc_3yr_xlsx": "Historical Counterparty Risk Graphs - ex LLC 3 Year.xlsx",
        "hist_llc_3yr_xlsx": "Historical Counterparty Risk Graphs - LLC 3 Year.xlsx",
        "monthly_pptx": "Monthly Counterparty Exposure Report.pptx",
    }
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    for name in names.values():
        shutil.copyfile(fixtures / name, inputs / name)
    cash = tmp_path / "repo-cash.csv"
    cash.write_text("counterparty,cash_value\nCIBC,1234567\n")
    return WorkflowConfig(
        **{field: inputs / name for field, name in names.items()},
        as_of_date=date(2025, 12, 31),
        cash_source_type="csv",
        cash_source_path=cash,
        enable_screenshot_replacement=True,
        output_root=tmp_path / "runs",
    )
