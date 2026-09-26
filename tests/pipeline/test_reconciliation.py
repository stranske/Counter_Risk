"""Focused reconciliation exception contract tests."""

from __future__ import annotations

from pathlib import Path

import pytest

from counter_risk.pipeline import run as run_module
from counter_risk.pipeline.parsing_types import UnmappedCounterpartyError
from counter_risk.pipeline.run import reconcile_series_coverage


def _cprs_ch_workbook(tmp_path: Path, sheet: str, label: str, notional: object) -> Path:
    from openpyxl import Workbook

    path = tmp_path / "totals.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet.cell(1, 3, "Unrelated row")
    worksheet.cell(3, 3, label)
    worksheet.cell(3, 11, notional)
    workbook.save(path)
    workbook.close()
    return path


def test_reconciliation_strict_exception_exposes_normalized_counterparty() -> None:
    try:
        reconcile_series_coverage(
            parsed_data_by_sheet={
                "Total": {"totals": [{"counterparty": "Bank of America, NA"}], "futures": []}
            },
            historical_series_headers_by_sheet={"Total": ("Legacy Counterparty",)},
            fail_policy="strict",
        )
    except UnmappedCounterpartyError as error:
        assert error.normalized_counterparty == "Bank of America"
    else:
        raise AssertionError("Expected UnmappedCounterpartyError")


def test_reconciliation_strict_exception_exposes_raw_counterparty() -> None:
    try:
        reconcile_series_coverage(
            parsed_data_by_sheet={
                "Total": {"totals": [{"counterparty": "Bank of America, NA"}], "futures": []}
            },
            historical_series_headers_by_sheet={"Total": ("Legacy Counterparty",)},
            fail_policy="strict",
        )
    except UnmappedCounterpartyError as error:
        assert error.raw_counterparty == "Bank of America, NA"
    else:
        raise AssertionError("Expected UnmappedCounterpartyError")


def test_reconciliation_warn_mode_records_structured_exception_without_raising() -> None:
    result = reconcile_series_coverage(
        parsed_data_by_sheet={
            "Total": {"totals": [{"counterparty": "Bank of America, NA"}], "futures": []}
        },
        historical_series_headers_by_sheet={"Total": ("Legacy Counterparty",)},
        fail_policy="warn",
    )

    exceptions = result.get("exceptions")
    assert isinstance(exceptions, list)
    assert len(exceptions) == 1
    assert isinstance(exceptions[0], UnmappedCounterpartyError)


def test_reconciliation_skips_unmanaged_sheet_when_present_sets_are_authoritative() -> None:
    result = reconcile_series_coverage(
        parsed_data_by_sheet={
            "Total": {"totals": [{"counterparty": "Counterparty A"}], "futures": []},
            "WAL": {"totals": [], "futures": []},
        },
        historical_series_headers_by_sheet={
            "Total": ("Counterparty A",),
            "WAL": ("Legacy Counterparty",),
        },
        series_present_by_sheet={"Total": ("Counterparty A",)},
    )

    assert set(result["by_sheet"]) == {"Total"}
    assert result["by_sheet"]["Total"]["counterparties_in_data"] == ["Counterparty A"]
    assert result["missing_series"] == []
    assert result["gap_count"] == 0


def test_reconciliation_prior_populated_treats_absent_series_as_dropped_gap() -> None:
    result = reconcile_series_coverage(
        parsed_data_by_sheet={
            "Total": {"totals": [], "futures": []},
        },
        historical_series_headers_by_sheet={
            "Total": ("Active Counterparty", "Dropped Counterparty", "Dormant Counterparty"),
        },
        prior_populated_series_by_sheet={
            "Total": ("Active Counterparty", "Dropped Counterparty"),
        },
    )

    sheet = result["by_sheet"]["Total"]
    assert sheet["dropped_from_data"] == ["Active Counterparty", "Dropped Counterparty"]
    assert sheet["dormant_from_data"] == ["Dormant Counterparty"]
    assert result["gap_count"] == 2
    assert any("missing from parsed data" in warning for warning in result["warnings"])


def test_reconciliation_prior_populated_classifies_dormant_headers_separately() -> None:
    result = reconcile_series_coverage(
        parsed_data_by_sheet={
            "Total": {
                "totals": [{"counterparty": "Active Counterparty"}],
                "futures": [],
            },
        },
        historical_series_headers_by_sheet={
            "Total": ("Active Counterparty", "Dormant Counterparty"),
        },
        prior_populated_series_by_sheet={
            "Total": ("Active Counterparty",),
        },
    )

    sheet = result["by_sheet"]["Total"]
    assert sheet["dropped_from_data"] == []
    assert sheet["dormant_from_data"] == ["Dormant Counterparty"]
    assert result["gap_count"] == 0
    assert any("dormant series" in warning.lower() for warning in result["warnings"])
    assert not any("missing from parsed data" in warning for warning in result["warnings"])


def test_reconciliation_missing_expected_segments_records_gap() -> None:
    result = reconcile_series_coverage(
        parsed_data_by_sheet={
            "Total": {
                "totals": [{"counterparty": "Counterparty A", "segment": "present"}],
                "futures": [],
            },
        },
        historical_series_headers_by_sheet={"Total": ("Counterparty A",)},
        variant="baseline",
        expected_segments_by_variant={"baseline": ("present", "missing-segment")},
    )

    sheet = result["by_sheet"]["Total"]
    assert sheet["missing_expected_segments"] == ["missing-segment"]
    assert result["missing_segments"] == [
        {
            "variant": "baseline",
            "sheet": "Total",
            "expected_segment_identifiers": ["missing-segment"],
        }
    ]
    assert result["gap_count"] == 1
    assert any("expected segments missing" in warning for warning in result["warnings"])


def test_reconcile_clearing_house_canonical_label_does_not_false_flag_missing_from_historical() -> (
    None
):
    result = reconcile_series_coverage(
        parsed_data_by_sheet={
            "Total": {
                "totals": [],
                "futures": [{"clearing_house": "ICE"}],
            },
        },
        historical_series_headers_by_sheet={"Total": ("ICE Clear US",)},
    )

    assert result["by_sheet"]["Total"]["missing_from_historical_headers"] == []
    assert result["gap_count"] == 0


@pytest.mark.parametrize("notional", [True, False, "nan", "inf", "-inf", "bad"])
def test_invalid_workbook_notional_fails_cprs_ch_totals_reconciliation(
    tmp_path: Path, notional: object
) -> None:
    path = _cprs_ch_workbook(tmp_path, "CPRS - CH", "MOSERS Program", notional)

    result = run_module._evaluate_cprs_ch_totals_reconciliation(
        parsed_sections={"totals": [{"Notional": 1.0}]},
        variant="all_programs",
        mosers_workbook_path=path,
    )

    assert result["status"] == "failed"
    assert "MOSERS Program row" in result["message"]
    assert "computed_total_notional" not in result


@pytest.mark.parametrize("sheet", [" CPRS - CH ", "Archived CPRS CH totals"])
@pytest.mark.parametrize("notional", [0, -125.5, "125.5"])
def test_cprs_ch_workbook_total_reads_finite_values(
    tmp_path: Path, sheet: str, notional: object
) -> None:
    path = _cprs_ch_workbook(tmp_path, sheet, " MOSERS PROGRAM ", notional)

    assert run_module._extract_mosers_program_notional_from_cprs_ch(workbook_path=path) == float(
        notional
    )


@pytest.mark.parametrize(
    ("sheet", "label", "notional"),
    [
        ("Unrelated", "MOSERS Program", 123),
        ("CPRS - CH", "Other program", 123),
        ("CPRS - CH", "MOSERS Program", None),
        ("CPRS - CH", "MOSERS Program", ""),
    ],
)
def test_absent_cprs_ch_workbook_total_returns_none(
    tmp_path: Path, sheet: str, label: str, notional: object
) -> None:
    path = _cprs_ch_workbook(tmp_path, sheet, label, notional)

    assert run_module._extract_mosers_program_notional_from_cprs_ch(workbook_path=path) is None
