"""Targeted PPT output tests for pipeline acceptance criteria."""

from __future__ import annotations

import shutil
from datetime import date
from pathlib import Path

import pytest
from pptx import Presentation

import counter_risk.pipeline.run as run_module
from counter_risk.config import WorkflowConfig
from counter_risk.outputs.pdf_export import PDFExportGenerator
from counter_risk.pipeline.manifest import ManifestBuilder
from counter_risk.pipeline.ppt_naming import resolve_ppt_output_names
from counter_risk.pipeline.ppt_validation import PptStandaloneValidationResult


def _write_placeholder(path: Path, *, payload: bytes = b"fixture") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)


def _build_config(
    tmp_path: Path, *, enable_ppt_output: bool, enable_distribution_output: bool = True
) -> WorkflowConfig:
    inputs_dir = tmp_path / "inputs"
    files = {
        "all_programs": inputs_dir / "all_programs.xlsx",
        "ex_trend": inputs_dir / "ex_trend.xlsx",
        "trend": inputs_dir / "trend.xlsx",
        "hist_all": inputs_dir / "hist_all.xlsx",
        "hist_ex": inputs_dir / "hist_ex.xlsx",
        "hist_llc": inputs_dir / "hist_llc.xlsx",
    }
    for input_path in files.values():
        _write_placeholder(input_path)

    return WorkflowConfig(
        mosers_all_programs_xlsx=files["all_programs"],
        mosers_ex_trend_xlsx=files["ex_trend"],
        mosers_trend_xlsx=files["trend"],
        hist_all_programs_3yr_xlsx=files["hist_all"],
        hist_ex_llc_3yr_xlsx=files["hist_ex"],
        hist_llc_3yr_xlsx=files["hist_llc"],
        monthly_pptx=Path("tests/fixtures/Monthly Counterparty Exposure Report.pptx"),
        output_root=tmp_path / "ignored-output-root",
        enable_screenshot_replacement=False,
        enable_ppt_output=enable_ppt_output,
        enable_distribution_output=enable_distribution_output,
    )


@pytest.mark.parametrize("explicit_outputs", [True, False])
@pytest.mark.parametrize("refresh_status", ["skipped", "success"])
def test_skipped_refresh_reports_existing_distribution(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    explicit_outputs: bool,
    refresh_status: str,
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    config = _build_config(tmp_path, enable_ppt_output=True)
    monkeypatch.setattr(
        run_module,
        "_refresh_ppt_links",
        lambda _path: run_module.PptProcessingResult(
            status=run_module.PptProcessingStatus(refresh_status),
        ),
    )
    warnings: list[str] = []
    output_paths, ppt_result = run_module._write_outputs(
        run_dir=run_dir, config=config, as_of_date=date(2025, 12, 31), warnings=warnings
    )
    distribution = run_dir / resolve_ppt_output_names(date(2025, 12, 31)).distribution_filename
    assert distribution in output_paths
    assert len(Presentation(str(distribution)).slides) > 0
    builder = ManifestBuilder(
        config=config, as_of_date=date(2025, 12, 31), run_date=date(2026, 1, 2)
    )
    manifest = builder.build(
        run_dir=run_dir,
        input_hashes={},
        output_paths=output_paths,
        top_exposures={},
        top_changes_per_variant={},
        warnings=warnings,
        ppt_status=ppt_result.status.value,
        ppt_outputs=ppt_result.ppt_outputs if explicit_outputs else None,
    )
    builder.write(run_dir=run_dir, manifest=manifest)
    summary = (run_dir / "DATA_QUALITY_SUMMARY.txt").read_text()
    codes = {finding["code"] for finding in manifest["data_quality"]["findings"]}
    assert manifest["ppt_outputs"]["distribution"]["status"] == "success"
    assert "PPT_GENERATION_SKIPPED" not in codes
    assert "PowerPoint generation was skipped" not in summary
    if refresh_status == "skipped":
        assert "PPT_LINK_REFRESH_SKIPPED" in codes
        assert "link refresh was skipped" in summary
        assert "Distribution PowerPoint was generated" in summary
        assert "WARN (YELLOW)" in summary
    else:
        assert "PPT_LINK_REFRESH_SKIPPED" not in codes


@pytest.mark.parametrize("enable_ppt_output", [True, False])
def test_skipped_refresh_reports_absent_distribution(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, enable_ppt_output: bool
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    config = _build_config(
        tmp_path, enable_ppt_output=enable_ppt_output, enable_distribution_output=False
    )
    monkeypatch.setattr(
        run_module,
        "_refresh_ppt_links",
        lambda _path: run_module.PptProcessingResult(status=run_module.PptProcessingStatus.SKIPPED),
    )
    output_paths, ppt_result = run_module._write_outputs(
        run_dir=run_dir, config=config, as_of_date=date(2025, 12, 31), warnings=[]
    )
    distribution = run_dir / resolve_ppt_output_names(date(2025, 12, 31)).distribution_filename
    assert not distribution.exists()
    builder = ManifestBuilder(
        config=config, as_of_date=date(2025, 12, 31), run_date=date(2026, 1, 2)
    )
    manifest = builder.build(
        run_dir=run_dir,
        input_hashes={},
        output_paths=output_paths,
        top_exposures={},
        top_changes_per_variant={},
        warnings=[],
        ppt_status=ppt_result.status.value,
        ppt_outputs=ppt_result.ppt_outputs,
    )
    builder.write(run_dir=run_dir, manifest=manifest)
    summary = (run_dir / "DATA_QUALITY_SUMMARY.txt").read_text()
    assert "PPT_GENERATION_SKIPPED" in summary
    assert "PowerPoint generation was skipped" in summary
    assert "Distribution PowerPoint was generated" not in summary
    assert ("PPT_LINK_REFRESH_SKIPPED" in summary) is enable_ppt_output


def test_ppt_disabled_skips_ppt_entrypoint_and_produces_no_pptx(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir(parents=True)
    config = _build_config(tmp_path, enable_ppt_output=False)

    def _unexpected_call(*args: object, **kwargs: object) -> None:
        _ = (args, kwargs)
        raise AssertionError("PPT code path should not be called when PPT output is disabled")

    monkeypatch.setattr(run_module, "_get_screenshot_replacer", _unexpected_call)
    monkeypatch.setattr(run_module, "_refresh_ppt_links", _unexpected_call)
    monkeypatch.setattr(run_module, "_derive_distribution_ppt", _unexpected_call)

    output_paths, ppt_result = run_module._write_outputs(
        run_dir=run_dir,
        config=config,
        as_of_date=date(2025, 12, 31),
        warnings=[],
    )

    assert ppt_result.status == run_module.PptProcessingStatus.SKIPPED
    assert all(path.suffix.lower() != ".pptx" for path in output_paths)
    assert list(run_dir.glob("*.pptx")) == []


def test_ppt_enabled_names_produce_exactly_two_expected_pptx_files(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir(parents=True)
    config = _build_config(tmp_path, enable_ppt_output=True)

    monkeypatch.setattr(
        run_module,
        "_refresh_ppt_links",
        lambda _path: run_module.PptProcessingResult(status=run_module.PptProcessingStatus.SUCCESS),
    )
    output_paths, ppt_result = run_module._write_outputs(
        run_dir=run_dir,
        config=config,
        as_of_date=date(2025, 12, 31),
        warnings=[],
    )

    assert sorted(path.name for path in run_dir.glob("*.pptx")) == [
        "Monthly Counterparty Exposure Report (Master) - 2025-12-31.pptx",
        "Monthly Counterparty Exposure Report - 2025-12-31.pptx",
    ]

    manifest = ManifestBuilder(
        config=config,
        as_of_date=date(2025, 12, 31),
        run_date=date(2026, 1, 2),
    ).build(
        run_dir=run_dir,
        input_hashes={},
        output_paths=output_paths,
        top_exposures={},
        top_changes_per_variant={},
        warnings=[],
        ppt_status=ppt_result.status.value,
    )
    assert manifest["ppt_outputs"] == {
        "distribution": {
            "generation_step": "ppt_distribution",
            "path": "Monthly Counterparty Exposure Report - 2025-12-31.pptx",
            "role": "distribution",
            "status": "success",
        },
        "master": {
            "generation_step": "ppt_master",
            "path": "Monthly Counterparty Exposure Report (Master) - 2025-12-31.pptx",
            "role": "maintainer_master",
            "status": "success",
        },
    }


def test_ppt_enabled_run_writes_readme_consistent_with_manifest_entries(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir(parents=True)
    config = _build_config(tmp_path, enable_ppt_output=True)
    as_of_date = date(2025, 12, 31)
    output_names = resolve_ppt_output_names(as_of_date)

    monkeypatch.setattr(
        run_module,
        "_refresh_ppt_links",
        lambda _path: run_module.PptProcessingResult(status=run_module.PptProcessingStatus.SUCCESS),
    )

    output_paths, ppt_result = run_module._write_outputs(
        run_dir=run_dir,
        config=config,
        as_of_date=as_of_date,
        warnings=[],
    )

    readme_path = run_dir / "README.txt"
    assert readme_path.exists()
    readme_content = readme_path.read_text(encoding="utf-8")
    assert output_names.master_filename in readme_content
    assert output_names.distribution_filename in readme_content

    manifest = ManifestBuilder(
        config=config,
        as_of_date=as_of_date,
        run_date=date(2026, 1, 2),
    ).build(
        run_dir=run_dir,
        input_hashes={},
        output_paths=output_paths,
        top_exposures={},
        top_changes_per_variant={},
        warnings=[],
        ppt_status=ppt_result.status.value,
    )

    assert "README.txt" in manifest["output_paths"]
    assert manifest["ppt_outputs"]["master"]["path"] == output_names.master_filename
    assert manifest["ppt_outputs"]["distribution"]["path"] == output_names.distribution_filename


def test_master_refresh_failure_logs_error_and_skips_distribution_derivation(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    caplog.set_level("ERROR")
    run_dir = tmp_path / "run"
    run_dir.mkdir(parents=True)
    config = _build_config(tmp_path, enable_ppt_output=True)
    calls = {"distribution": 0}

    def _refresh_raises(_path: Path) -> run_module.PptProcessingResult:
        raise RuntimeError("refresh exploded")

    def _derive_distribution(*args: object, **kwargs: object) -> None:
        _ = (args, kwargs)
        calls["distribution"] += 1

    monkeypatch.setattr(run_module, "_refresh_ppt_links", _refresh_raises)
    monkeypatch.setattr(run_module, "_derive_distribution_ppt", _derive_distribution)

    output_paths, ppt_result = run_module._write_outputs(
        run_dir=run_dir,
        config=config,
        as_of_date=date(2025, 12, 31),
        warnings=[],
    )

    assert ppt_result.status == run_module.PptProcessingStatus.FAILED
    assert "refresh exploded" in (ppt_result.error_detail or "")
    assert calls["distribution"] == 0
    assert any(
        "Master" in record.message
        and "PPT" in record.message
        and "refresh exploded" in record.message
        for record in caplog.records
        if record.levelname == "ERROR"
    )
    assert all(
        path.name != "Monthly Counterparty Exposure Report - 2025-12-31.pptx"
        for path in output_paths
    )


def test_ppt_enabled_order_master_generated_before_distribution(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir(parents=True)
    config = _build_config(tmp_path, enable_ppt_output=True)
    config.enable_screenshot_replacement = True

    call_order: list[str] = []

    def _replace_master(source: Path, target: Path, screenshot_inputs: dict[str, Path]) -> None:
        _ = screenshot_inputs
        call_order.append("master_generation")
        shutil.copy2(source, target)

    def _refresh_links(_path: Path) -> run_module.PptProcessingResult:
        call_order.append("master_refresh")
        return run_module.PptProcessingResult(status=run_module.PptProcessingStatus.SUCCESS)

    def _derive_distribution(*, master_pptx_path: Path, distribution_pptx_path: Path) -> None:
        call_order.append("distribution_derivation")
        shutil.copy2(master_pptx_path, distribution_pptx_path)

    monkeypatch.setattr(run_module, "_resolve_screenshot_input_mapping", lambda _config: {})
    monkeypatch.setattr(run_module, "_get_screenshot_replacer", lambda _impl: _replace_master)
    monkeypatch.setattr(run_module, "_refresh_ppt_links", _refresh_links)
    monkeypatch.setattr(run_module, "_derive_distribution_ppt", _derive_distribution)
    monkeypatch.setattr(
        run_module,
        "validate_distribution_ppt_standalone",
        lambda _path: PptStandaloneValidationResult(
            is_valid=True,
            external_relationship_count=0,
            relationship_parts_scanned=(),
            external_relationship_parts=(),
        ),
    )

    run_module._write_outputs(
        run_dir=run_dir,
        config=config,
        as_of_date=date(2025, 12, 31),
        warnings=[],
    )

    assert call_order.index("master_generation") < call_order.index("distribution_derivation")
    assert call_order.count("distribution_derivation") == 1


def test_distribution_disabled_keeps_master_and_records_skipped_manifest_entry(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir(parents=True)
    config = _build_config(
        tmp_path,
        enable_ppt_output=True,
        enable_distribution_output=False,
    )
    as_of_date = date(2025, 12, 31)
    output_names = resolve_ppt_output_names(as_of_date)
    calls = {"distribution": 0}

    monkeypatch.setattr(
        run_module,
        "_refresh_ppt_links",
        lambda _path: run_module.PptProcessingResult(status=run_module.PptProcessingStatus.SUCCESS),
    )

    def _derive_distribution(*args: object, **kwargs: object) -> None:
        _ = (args, kwargs)
        calls["distribution"] += 1

    monkeypatch.setattr(run_module, "_derive_distribution_ppt", _derive_distribution)

    output_paths, ppt_result = run_module._write_outputs(
        run_dir=run_dir,
        config=config,
        as_of_date=as_of_date,
        warnings=[],
    )

    assert (run_dir / output_names.master_filename).exists()
    assert not (run_dir / output_names.distribution_filename).exists()
    assert not (run_dir / "README.txt").exists()
    assert calls["distribution"] == 0
    assert all(Path(path).name != output_names.distribution_filename for path in output_paths)

    manifest = ManifestBuilder(
        config=config,
        as_of_date=as_of_date,
        run_date=date(2026, 1, 2),
    ).build(
        run_dir=run_dir,
        input_hashes={},
        output_paths=output_paths,
        top_exposures={},
        top_changes_per_variant={},
        warnings=[],
        ppt_status=ppt_result.status.value,
    )

    assert manifest["ppt_outputs"]["master"]["status"] == "success"
    assert manifest["ppt_outputs"]["distribution"] == {
        "generation_step": "ppt_distribution",
        "path": output_names.distribution_filename,
        "role": "distribution",
        "skipped_reason": "Distribution output disabled for this run.",
        "status": "skipped",
    }


def test_pdf_request_with_distribution_disabled_is_explicit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    exported: list[tuple[Path, Path]] = []

    def _export_pdf(source: Path, target: Path) -> None:
        exported.append((source, target))
        target.write_bytes(b"%PDF-1.4\n%test\n")

    monkeypatch.setattr(
        run_module,
        "_refresh_ppt_links",
        lambda _path: run_module.PptProcessingResult(status=run_module.PptProcessingStatus.SUCCESS),
    )
    monkeypatch.setattr(
        run_module,
        "_build_pdf_export_output_generator",
        lambda *, source_pptx, warnings: PDFExportGenerator(
            source_pptx=source_pptx,
            warnings=warnings,
            com_availability_checker=lambda: True,
            pptx_to_pdf_exporter=_export_pdf,
        ),
    )

    disabled_config = _build_config(
        tmp_path / "disabled", enable_ppt_output=True, enable_distribution_output=False
    )
    disabled_config.export_pdf = True
    disabled_dir = tmp_path / "disabled" / "run"
    disabled_dir.mkdir(parents=True)
    with pytest.raises(ValueError, match="export_pdf requires enable_distribution_output"):
        run_module._validate_pipeline_config(disabled_config)
    with pytest.raises(ValueError, match="export_pdf requires enable_distribution_output"):
        run_module._write_outputs(
            run_dir=disabled_dir,
            config=disabled_config,
            as_of_date=date(2025, 12, 31),
            warnings=[],
        )
    assert exported == []
    assert list(disabled_dir.iterdir()) == []

    enabled_config = _build_config(tmp_path / "enabled", enable_ppt_output=True)
    enabled_config.export_pdf = True
    run_module._validate_pipeline_config(enabled_config)
    enabled_dir = tmp_path / "enabled" / "run"
    enabled_dir.mkdir(parents=True)
    output_paths, _ = run_module._write_outputs(
        run_dir=enabled_dir,
        config=enabled_config,
        as_of_date=date(2025, 12, 31),
        warnings=[],
    )
    assert len(exported) == 1
    source, target = exported[0]
    assert (
        source == enabled_dir / resolve_ppt_output_names(date(2025, 12, 31)).distribution_filename
    )
    assert target == source.with_suffix(".pdf")
    assert target in output_paths
    assert target.read_bytes().startswith(b"%PDF-1.4")

    no_pdf_config = _build_config(
        tmp_path / "no-pdf", enable_ppt_output=True, enable_distribution_output=True
    )
    no_pdf_config.export_pdf = False
    run_module._validate_pipeline_config(no_pdf_config)
    no_pdf_dir = tmp_path / "no-pdf" / "run"
    no_pdf_dir.mkdir(parents=True)
    output_paths, _ = run_module._write_outputs(
        run_dir=no_pdf_dir,
        config=no_pdf_config,
        as_of_date=date(2025, 12, 31),
        warnings=[],
    )
    assert not any(path.suffix == ".pdf" for path in output_paths)
    assert len(exported) == 1


@pytest.mark.parametrize(
    ("include_concentration", "export_pdf", "expected_slides"),
    [(True, True, 24), (False, True, 23), (True, False, 24)],
)
def test_final_concentration_slide_precedes_pdf_export(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    include_concentration: bool,
    export_pdf: bool,
    expected_slides: int,
) -> None:
    fixtures = Path("tests/fixtures")
    config_path = tmp_path / "config.yml"
    config_path.write_text(
        "\n".join(
            [
                "as_of_date: 2025-12-31",
                f"mosers_all_programs_xlsx: {fixtures / 'MOSERS Counterparty Risk Summary 12-31-2025 - All Programs.xlsx'}",
                f"mosers_ex_trend_xlsx: {fixtures / 'MOSERS Counterparty Risk Summary 12-31-2025 - Ex Trend.xlsx'}",
                f"mosers_trend_xlsx: {fixtures / 'MOSERS Counterparty Risk Summary 12-31-2025 - Trend.xlsx'}",
                f"hist_all_programs_3yr_xlsx: {fixtures / 'Historical Counterparty Risk Graphs - All Programs 3 Year.xlsx'}",
                f"hist_ex_llc_3yr_xlsx: {fixtures / 'Historical Counterparty Risk Graphs - ex LLC 3 Year.xlsx'}",
                f"hist_llc_3yr_xlsx: {fixtures / 'Historical Counterparty Risk Graphs - LLC 3 Year.xlsx'}",
                f"monthly_pptx: {fixtures / 'Monthly Counterparty Exposure Report.pptx'}",
                f"output_root: {tmp_path / 'runs'}",
                f"include_concentration_table_in_ppt: {str(include_concentration).lower()}",
                f"export_pdf: {str(export_pdf).lower()}",
            ]
        )
        + "\n",
        encoding="utf-8",
    )
    export_source: list[tuple[int, list[str]]] = []
    validated_slides: list[int] = []
    real_validate = run_module.validate_distribution_ppt_standalone

    def _last_table_row(pptx_path: Path) -> list[str]:
        presentation = Presentation(str(pptx_path))
        tables = [shape.table for shape in presentation.slides[-1].shapes if shape.has_table]
        if not tables:
            return []
        return [cell.text for cell in tables[0].rows[1].cells]

    def _observe_validation(pptx_path: Path) -> PptStandaloneValidationResult:
        validated_slides.append(len(Presentation(str(pptx_path)).slides))
        return real_validate(pptx_path)

    def _export_pdf(source: Path, target: Path) -> None:
        export_source.append((len(Presentation(str(source)).slides), _last_table_row(source)))
        target.write_bytes(b"%PDF-1.4\n%test\n")

    monkeypatch.setattr(run_module, "validate_distribution_ppt_standalone", _observe_validation)
    monkeypatch.setattr(
        run_module,
        "_build_pdf_export_output_generator",
        lambda *, source_pptx, warnings: PDFExportGenerator(
            source_pptx=source_pptx,
            warnings=warnings,
            com_availability_checker=lambda: True,
            pptx_to_pdf_exporter=_export_pdf,
        ),
    )

    run_dir = run_module.run_pipeline(config_path, output_dir=tmp_path / "run")
    distribution = run_dir / resolve_ppt_output_names(date(2025, 12, 31)).distribution_filename
    assert len(Presentation(str(distribution)).slides) == expected_slides
    assert validated_slides == [expected_slides]
    final_row = _last_table_row(distribution)
    if include_concentration:
        assert len(final_row) == 5
        assert final_row[0] == "all_programs"
    else:
        assert final_row == []
    if export_pdf:
        assert export_source == [(expected_slides, final_row)]
        assert distribution.with_suffix(".pdf").exists()
    else:
        assert export_source == []
        assert not distribution.with_suffix(".pdf").exists()


def test_final_concentration_validation_failure_prevents_pdf_export(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    config = _build_config(tmp_path, enable_ppt_output=True)
    config.include_concentration_table_in_ppt = True
    config.export_pdf = True
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    exported: list[Path] = []

    def _invalid_final_deck(path: Path) -> PptStandaloneValidationResult:
        assert len(Presentation(str(path)).slides) == 24
        return PptStandaloneValidationResult(
            is_valid=False,
            external_relationship_count=1,
            relationship_parts_scanned=("ppt/slides/_rels/slide24.xml.rels",),
            external_relationship_parts=("ppt/slides/_rels/slide24.xml.rels",),
        )

    def _unexpected_export(source: Path, target: Path) -> None:
        _ = target
        exported.append(source)

    monkeypatch.setattr(run_module, "validate_distribution_ppt_standalone", _invalid_final_deck)
    monkeypatch.setattr(
        run_module,
        "_build_pdf_export_output_generator",
        lambda *, source_pptx, warnings: PDFExportGenerator(
            source_pptx=source_pptx,
            warnings=warnings,
            com_availability_checker=lambda: True,
            pptx_to_pdf_exporter=_unexpected_export,
        ),
    )

    with pytest.raises(RuntimeError, match="standalone validation failed"):
        run_module._write_outputs(
            run_dir=run_dir,
            config=config,
            as_of_date=date(2025, 12, 31),
            warnings=[],
            concentration_metrics_records=[{"variant": "all_programs", "segment": "fixture"}],
        )
    assert exported == []


def test_master_refresh_skipped_records_master_skipped_and_distribution_success(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir(parents=True)
    config = _build_config(tmp_path, enable_ppt_output=True)
    as_of_date = date(2025, 12, 31)
    output_names = resolve_ppt_output_names(as_of_date)

    monkeypatch.setattr(
        run_module,
        "_refresh_ppt_links",
        lambda _path: run_module.PptProcessingResult(
            status=run_module.PptProcessingStatus.SKIPPED,
            error_detail="unsupported platform",
        ),
    )

    output_paths, ppt_result = run_module._write_outputs(
        run_dir=run_dir,
        config=config,
        as_of_date=as_of_date,
        warnings=[],
    )

    manifest = ManifestBuilder(
        config=config,
        as_of_date=as_of_date,
        run_date=date(2026, 1, 2),
    ).build(
        run_dir=run_dir,
        input_hashes={},
        output_paths=output_paths,
        top_exposures={},
        top_changes_per_variant={},
        warnings=[],
        ppt_status=ppt_result.status.value,
        ppt_outputs=ppt_result.ppt_outputs,
    )

    assert manifest["ppt_outputs"]["master"] == {
        "generation_step": "ppt_master",
        "path": output_names.master_filename,
        "role": "maintainer_master",
        "skipped_reason": "unsupported platform",
        "status": "skipped",
    }
    assert manifest["ppt_outputs"]["distribution"] == {
        "generation_step": "ppt_distribution",
        "path": output_names.distribution_filename,
        "role": "distribution",
        "status": "success",
    }


def test_no_distribution_without_master_when_master_generation_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    run_dir = tmp_path / "run"
    run_dir.mkdir(parents=True)
    config = _build_config(tmp_path, enable_ppt_output=True)
    as_of_date = date(2025, 12, 31)
    distribution_name = f"Monthly Counterparty Exposure Report - {as_of_date.isoformat()}.pptx"

    def _refresh_raises(_path: Path) -> run_module.PptProcessingResult:
        raise RuntimeError("refresh exploded")

    monkeypatch.setattr(run_module, "_refresh_ppt_links", _refresh_raises)

    output_paths, _ppt_result = run_module._write_outputs(
        run_dir=run_dir,
        config=config,
        as_of_date=as_of_date,
        warnings=[],
    )

    assert not (run_dir / distribution_name).exists()
    assert not (run_dir / "README.txt").exists()

    manifest = ManifestBuilder(
        config=config,
        as_of_date=as_of_date,
        run_date=date(2026, 1, 2),
    ).build(
        run_dir=run_dir,
        input_hashes={},
        output_paths=output_paths,
        top_exposures={},
        top_changes_per_variant={},
        warnings=[],
        ppt_status=run_module.PptProcessingStatus.FAILED.value,
    )
    assert all(Path(path).name != distribution_name for path in manifest["output_paths"])
    assert all(Path(path).name != "README.txt" for path in manifest["output_paths"])
    assert "distribution" not in manifest.get("ppt_outputs", {})
    assert manifest["ppt_outputs"]["master"] == {
        "generation_step": "ppt_master",
        "path": "Monthly Counterparty Exposure Report (Master) - 2025-12-31.pptx",
        "role": "maintainer_master",
        "status": "failed",
    }
