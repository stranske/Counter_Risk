"""Exercise the workbook fallback without production-sized workbook fixtures."""

from pathlib import Path

import pytest

from counter_risk.pipeline import run as run_module


def _workbook(tmp_path: Path, sheet: str, label: str, notional: object) -> Path:
    from openpyxl import Workbook

    path = tmp_path / "totals.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet.cell(1, 3, "Unrelated row")
    worksheet.cell(3, 3, label)
    worksheet.cell(3, 11, notional)
    workbook.save(path)
    workbook.close()
    return path


@pytest.mark.parametrize("notional", [True, False, "nan", "inf", "-inf", "bad"])
def test_invalid_workbook_notional_fails_reconciliation(tmp_path: Path, notional: object) -> None:
    path = _workbook(tmp_path, "CPRS - CH", "MOSERS Program", notional)

    result = run_module._evaluate_cprs_ch_totals_reconciliation(
        parsed_sections={"totals": [{"Notional": 1.0}]},
        variant="all_programs",
        mosers_workbook_path=path,
    )

    assert result["status"] == "failed"
    assert "MOSERS Program row" in result["message"]
    assert "computed_total_notional" not in result


@pytest.mark.parametrize("sheet", [" CPRS - CH ", "Archived CPRS CH totals"])
@pytest.mark.parametrize("notional", [0, -125.5, "125.5"])
def test_workbook_total_reads_finite_values(tmp_path: Path, sheet: str, notional: object) -> None:
    path = _workbook(tmp_path, sheet, " MOSERS PROGRAM ", notional)

    assert run_module._extract_mosers_program_notional_from_cprs_ch(workbook_path=path) == float(
        notional
    )


@pytest.mark.parametrize(
    ("sheet", "label", "notional"),
    [
        ("Unrelated", "MOSERS Program", 123),
        ("CPRS - CH", "Other program", 123),
        ("CPRS - CH", "MOSERS Program", None),
        ("CPRS - CH", "MOSERS Program", ""),
    ],
)
def test_absent_workbook_total_returns_none(
    tmp_path: Path, sheet: str, label: str, notional: object
) -> None:
    path = _workbook(tmp_path, sheet, label, notional)

    assert run_module._extract_mosers_program_notional_from_cprs_ch(workbook_path=path) is None
