"""Workbook headers remain bounded and tolerate unavailable input without leaking readers."""

from __future__ import annotations

import builtins
from pathlib import Path
from zipfile import ZipFile

import pytest

from counter_risk.pipeline import run as run_module


def _workbook(path: Path) -> None:
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["  CPRS  ", None, "  ", 42, "outside columns"])
    sheet.append([None, None, None, None, "also outside"])
    sheet.append(["outside rows"])
    book.save(path)
    book.close()


def _track_reader(monkeypatch: pytest.MonkeyPatch, path: Path):
    import openpyxl

    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    closed = []
    original_close = book.close

    def close():
        closed.append(True)
        original_close()

    monkeypatch.setattr(book, "close", close)
    monkeypatch.setattr(openpyxl, "load_workbook", lambda **kwargs: book)
    return book, closed


def test_header_extraction_bounds_rows_columns_and_closes_reader(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "headers.xlsx"
    _workbook(path)
    book, closed = _track_reader(monkeypatch, path)

    assert run_module._extract_header_text_lines(path, max_rows=2, max_cols=4) == ["CPRS 42"]
    assert closed == [True]
    assert book._archive.fp is None


@pytest.mark.parametrize("error", [ModuleNotFoundError, RuntimeError], ids=["missing", "broken"])
def test_header_extraction_tolerates_unavailable_optional_reader(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, error: type[Exception]
) -> None:
    original_import = builtins.__import__

    def guarded_import(name, *args, **kwargs):
        if name == "openpyxl":
            raise error("optional workbook reader unavailable")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", guarded_import)

    assert run_module._extract_header_text_lines(tmp_path / "unused.xlsx") == []


def test_header_extraction_tolerates_corrupt_workbook(tmp_path: Path) -> None:
    path = tmp_path / "corrupt.xlsx"
    original_bytes = b"This is not an XLSX zip archive."
    path.write_bytes(original_bytes)

    assert run_module._extract_header_text_lines(path) == []
    assert path.read_bytes() == original_bytes


def test_header_extraction_handles_no_worksheets_and_closes_reader(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "source.xlsx"
    _workbook(source)
    path = tmp_path / "empty.xlsx"
    # A real XLSX with an empty sheet inventory, rather than a stubbed workbook.
    with ZipFile(source) as archive, ZipFile(path, "w") as empty:
        for entry in archive.infolist():
            data = archive.read(entry.filename)
            if entry.filename == "xl/workbook.xml":
                from xml.etree import ElementTree

                root = ElementTree.fromstring(data)
                sheets = root.find(
                    "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheets"
                )
                assert sheets is not None
                sheets.clear()
                data = ElementTree.tostring(root)
            empty.writestr(entry, data)
    book, closed = _track_reader(monkeypatch, path)
    assert book.sheetnames == []

    assert run_module._extract_header_text_lines(path) == []
    assert closed == [True]
    assert book._archive.fp is None


def test_header_extraction_discards_partial_read_and_closes_reader(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "unreadable-row.xlsx"
    _workbook(path)
    book, closed = _track_reader(monkeypatch, path)
    sheet = book[book.sheetnames[0]]
    original_cell = sheet.cell

    def cell(*, row, column):
        if row == 2:
            raise OSError("workbook row became unreadable")
        return original_cell(row=row, column=column)

    monkeypatch.setattr(sheet, "cell", cell)

    assert run_module._extract_header_text_lines(path) == []
    assert closed == [True]
    assert book._archive.fp is None
