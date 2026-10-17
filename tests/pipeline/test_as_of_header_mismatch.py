"""Date safeguard orchestration with real bounded XLSX reads and durable artifacts."""

from __future__ import annotations

import hashlib
import json
from datetime import date

import pytest
import yaml

from counter_risk.cli import main
from counter_risk.config import WorkflowConfig
from counter_risk.pipeline import run as pipeline

Workbook = pytest.importorskip("openpyxl", reason="XLSX date safeguards require openpyxl").Workbook


@pytest.fixture
def date_run(tmp_path, monkeypatch):
    """Isolate unrelated calculations; run actual date, config, manifest and CLI boundaries."""
    paths = {}
    for variant in ("all_programs", "ex_trend", "trend"):
        path = tmp_path / f"{variant}.xlsx"
        book = Workbook()
        book.active.append(["Report Date:", "2026-01-31"])
        book.save(path)
        book.close()
        paths[f"mosers_{variant}_xlsx"] = path
    for field in (
        "hist_all_programs_3yr_xlsx",
        "hist_ex_llc_3yr_xlsx",
        "hist_llc_3yr_xlsx",
        "monthly_pptx",
    ):
        path = tmp_path / (field + (".pptx" if field == "monthly_pptx" else ".xlsx"))
        path.write_bytes(b"original history/publication input")
        paths[field] = path
    config = WorkflowConfig(
        as_of_date=date(2026, 1, 31),
        run_date=date(2026, 3, 1),
        output_root=tmp_path / "runs",
        enable_ppt_output=False,
        enable_distribution_output=False,
        output_generators=(),
        **paths,
    )
    calls = []
    monkeypatch.setattr(pipeline, "_resolve_repo_root", lambda: tmp_path)
    monkeypatch.setattr(pipeline, "_parse_inputs", lambda *a: {})
    monkeypatch.setattr(pipeline, "_validate_parsed_inputs", lambda *a: None)
    monkeypatch.setattr(pipeline, "_run_reconciliation_checks", lambda **k: {})
    monkeypatch.setattr(pipeline, "_compute_metrics", lambda *a: ({}, {}))
    monkeypatch.setattr(pipeline, "_write_risk_outputs", lambda **k: [])
    monkeypatch.setattr(pipeline, "_compute_and_write_concentration_metrics", lambda **k: [])
    monkeypatch.setattr(pipeline, "_write_change_attribution_outputs", lambda **k: [])
    monkeypatch.setattr(
        pipeline,
        "_compute_and_write_limit_breaches",
        lambda **k: pipeline.LimitBreachEvaluation(None, 0),
    )

    def history(**kwargs):
        calls.append("history")
        return []

    def publication(**kwargs):
        calls.append("publication")
        return []

    monkeypatch.setattr(pipeline, "_update_historical_outputs", history)
    monkeypatch.setattr(pipeline, "_call_write_outputs", publication)
    return config, calls


def _write_config(config, path):
    path.write_text(yaml.safe_dump(config.model_dump(mode="json")))
    return path


def _set_header(path, lines):
    book = Workbook()
    for line in lines:
        book.active.append([line])
    book.save(path)
    book.close()


def _manifest(config, tmp_path, **kwargs):
    run_dir = pipeline.run_pipeline(_write_config(config, tmp_path / "config.yml"), **kwargs)
    manifest = json.loads((run_dir / "manifest.json").read_text())
    return run_dir, manifest, manifest["date_resolution"]["source_report_dates"]


@pytest.mark.parametrize("configured", [date(2026, 1, 31), None])
def test_later_variant_conflict_retains_resolution_and_source_details(
    date_run, tmp_path, configured
):
    config, _ = date_run
    config = config.model_copy(update={"as_of_date": configured})
    _set_header(config.mosers_trend_xlsx, ["Report Date: 2026-02-28"])
    run_dir, manifest, comparison = _manifest(config, tmp_path)
    assert manifest["as_of_date"] == "2026-01-31"
    assert manifest["run_date"] == "2026-03-01"
    assert comparison["inter_variant_conflict"] is True
    assert comparison["config_conflict_variants"] == (["trend"] if configured else [])
    finding = next(
        f
        for f in manifest["data_quality"]["findings"]
        if f["code"] == "SOURCE_REPORT_DATE_CONFLICT"
    )
    assert finding["severity"] == "warn"
    assert finding["details"] == comparison
    source = finding["details"]["sources"][2]
    assert source["variant"] == "trend"
    assert source["path"] == str(config.mosers_trend_xlsx)
    assert source["report_date"] == "2026-02-28"
    assert "2026-01-31" in finding["message"] and "2026-02-28" in finding["message"]
    assert manifest["data_quality"]["overall_status"] == "warn"
    assert "YELLOW" in (run_dir / "DATA_QUALITY_SUMMARY.txt").read_text()


def test_first_variant_config_mismatch_is_advisory(date_run, tmp_path):
    config, calls = date_run
    _set_header(config.mosers_all_programs_xlsx, ["As of Date: 2026-02-28"])
    _, manifest, comparison = _manifest(config, tmp_path)
    assert manifest["as_of_date"] == "2026-01-31"
    assert comparison["config_conflict_variants"] == ["all_programs"]
    assert calls == ["history", "publication"]  # disabled stages are safe no-ops in diagnostics


def test_matching_sources_have_no_conflict_finding(date_run, tmp_path):
    config, _ = date_run
    _, manifest, comparison = _manifest(config, tmp_path)
    assert comparison["status"] == "consistent"
    assert comparison["has_conflict"] is False
    assert all(
        f["code"] != "SOURCE_REPORT_DATE_CONFLICT" for f in manifest["data_quality"]["findings"]
    )


@pytest.mark.parametrize(
    "mode", ["missing", "unreadable", "ambiguous", "generated", "legacy_generated"]
)
def test_unverified_workbook_never_confirms_match_or_conflict(date_run, tmp_path, mode):
    config, _ = date_run
    path = config.mosers_trend_xlsx
    if mode == "unreadable":
        path.write_bytes(b"corrupt XLSX")
    else:
        _set_header(
            path,
            {
                "missing": ["Created 2026-02-28"],
                "ambiguous": ["Report Date: 2026-01-31", "Report Date: 2026-02-28"],
                "generated": ["Report Date: 2026-02-28", "Generated from config as_of_date"],
                "legacy_generated": ["CPRS CH Header Date 2026-02-28"],
            }[mode],
        )
    _, _, comparison = _manifest(config, tmp_path)
    assert comparison["status"] == "unverified"
    assert comparison["has_conflict"] is False
    assert comparison["sources"][2]["report_date"] is None


@pytest.mark.parametrize("mode", ["strict", "history", "publication", "distribution", "pdf"])
def test_strict_or_final_conflict_exits_with_diagnostics_before_side_effects(
    date_run, tmp_path, mode
):
    config, calls = date_run
    updates = {
        "strict": {"report_date_policy": "strict"},
        "history": {
            "output_generators": WorkflowConfig.model_fields["output_generators"].default_factory()
        },
        "publication": {"enable_ppt_output": True},
        "distribution": {"enable_distribution_output": True},
        "pdf": {"export_pdf": True, "enable_distribution_output": True},
    }[mode]
    config = config.model_copy(update=updates)
    _set_header(config.mosers_trend_xlsx, ["Report Date: 2026-02-28"])
    original = {
        path: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in (
            config.hist_all_programs_3yr_xlsx,
            config.hist_ex_llc_3yr_xlsx,
            config.hist_llc_3yr_xlsx,
            config.monthly_pptx,
        )
    }
    path = _write_config(config, tmp_path / "config.yml")
    assert main(["run", "--config", str(path)]) == 1
    assert calls == []
    run_dir = next(config.output_root.iterdir())
    manifest = json.loads((run_dir / "manifest.json").read_text())
    assert manifest["date_resolution"]["source_report_dates"]["execution"] == "blocked"
    assert "SOURCE_REPORT_DATE_CONFLICT" in (run_dir / "DATA_QUALITY_SUMMARY.txt").read_text()
    assert {p.name for p in run_dir.iterdir()} == {"manifest.json", "DATA_QUALITY_SUMMARY.txt"}
    for path, digest in original.items():
        assert hashlib.sha256(path.read_bytes()).hexdigest() == digest


def test_run_level_override_is_audited_and_reaches_final_stages(date_run, tmp_path):
    config, calls = date_run
    config = config.model_copy(update={"report_date_policy": "strict", "enable_ppt_output": True})
    _set_header(config.mosers_trend_xlsx, ["Report Date: 2026-02-28"])
    path = _write_config(config, tmp_path / "config.yml")
    assert (
        main(
            [
                "run",
                "--config",
                str(path),
                "--report-date-override",
                "Risk owner approved period exception REF-42",
            ]
        )
        == 0
    )
    assert calls == ["history", "publication"]
    run_dir = next(config.output_root.iterdir())
    manifest = json.loads((run_dir / "manifest.json").read_text())
    comparison = manifest["date_resolution"]["source_report_dates"]
    assert comparison["execution"] == "override"
    assert comparison["override"]["reason"] == "Risk owner approved period exception REF-42"
    assert comparison["override"]["run_dir"] == str(run_dir)
    assert comparison["override"]["run_date"] == "2026-03-01"
    assert "report_date_override" not in manifest["config_snapshot"]
    assert "YELLOW" in (run_dir / "DATA_QUALITY_SUMMARY.txt").read_text()
    # Override is not retained by the next invocation of the same config.
    assert main(["run", "--config", str(path)]) == 1
    assert calls == ["history", "publication"]


def test_blank_override_cannot_bypass_conflict(date_run, tmp_path, capsys):
    config, calls = date_run
    _set_header(config.mosers_trend_xlsx, ["Report Date: 2026-02-28"])
    assert (
        main(
            [
                "run",
                "--config",
                str(_write_config(config, tmp_path / "config.yml")),
                "--report-date-override",
                "   ",
            ]
        )
        == 1
    )
    assert calls == []
    assert "non-empty audit reason" in capsys.readouterr().out


def test_real_pipeline_parsers_and_outputs_keep_date_conflict_diagnostics(tmp_path):
    """Exercise U05 through the real pipeline without replacing calculation/output stages."""
    from pathlib import Path

    from openpyxl import load_workbook

    fixtures = Path("tests/fixtures").resolve()
    original = fixtures / "MOSERS Counterparty Risk Summary 12-31-2025 - All Programs.xlsx"
    source = tmp_path / "all.xlsx"
    book = load_workbook(original)
    book[book.sheetnames[0]]["A1"] = "Report Date: 2026-02-28"
    book.save(source)
    book.close()
    config = WorkflowConfig(
        as_of_date=date(2026, 1, 31),
        run_date=date(2026, 3, 1),
        mosers_all_programs_xlsx=source,
        mosers_ex_trend_xlsx=fixtures
        / "MOSERS Counterparty Risk Summary 12-31-2025 - Ex Trend.xlsx",
        mosers_trend_xlsx=fixtures / "MOSERS Counterparty Risk Summary 12-31-2025 - Trend.xlsx",
        hist_all_programs_3yr_xlsx=fixtures
        / "Historical Counterparty Risk Graphs - All Programs 3 Year.xlsx",
        hist_ex_llc_3yr_xlsx=fixtures / "Historical Counterparty Risk Graphs - ex LLC 3 Year.xlsx",
        hist_llc_3yr_xlsx=fixtures / "Historical Counterparty Risk Graphs - LLC 3 Year.xlsx",
        monthly_pptx=fixtures / "Monthly Counterparty Exposure Report.pptx",
        enable_ppt_output=False,
        enable_distribution_output=False,
        output_generators=(),
    )
    run_dir, manifest, comparison = _manifest(config, tmp_path, output_dir=tmp_path / "diagnostics")
    assert manifest["as_of_date"] == "2026-01-31"
    assert comparison["config_conflict_variants"] == ["all_programs"]
    assert manifest["top_exposures"]["all_programs"]  # real parser/calculation path
    assert any(
        f["code"] == "SOURCE_REPORT_DATE_CONFLICT" for f in manifest["data_quality"]["findings"]
    )
    assert "2026-02-28" in (run_dir / "DATA_QUALITY_SUMMARY.txt").read_text()


def test_named_date_sheet_is_bounded_and_generated_stamp_is_unverified(date_run):
    from counter_risk.writers.mosers_workbook import _stamp_as_of_date_on_workbook

    config, _ = date_run
    book = Workbook()
    book.active.title = "Cover"
    sheet = book.create_sheet("CPRS - CH")
    sheet["A2"] = "Report Date: 2026-02-28"
    sheet["A16"] = "Report Date: 2026-03-31"  # beyond row boundary
    sheet["G4"] = "Report Date: 2026-04-30"  # beyond column boundary
    book.create_sheet("Unrelated")["A1"] = "Report Date: 2026-05-31"
    book.save(config.mosers_all_programs_xlsx)
    evidence = pipeline._collect_source_report_dates(config=config)[0]
    assert evidence["report_date"] == "2026-02-28"
    _stamp_as_of_date_on_workbook(book, date(2026, 1, 31))
    book.save(config.mosers_all_programs_xlsx)
    book.close()
    evidence = pipeline._collect_source_report_dates(config=config)[0]
    assert evidence["status"] == "generated"
    assert evidence["report_date"] is None


def test_raw_replacement_cannot_confirm_its_config_date(date_run):
    config, _ = date_run
    config = config.model_copy(
        update={"raw_nisa_all_programs_xlsx": config.mosers_all_programs_xlsx}
    )
    evidence = pipeline._collect_source_report_dates(config=config)[0]
    assert evidence["status"] == "generated"
    assert evidence["report_date"] is None
