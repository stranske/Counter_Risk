"""Focused pipeline orchestration coverage for coverage-autopilot issue #1214."""

from __future__ import annotations

import sys
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, Mock

import pytest

from counter_risk.pipeline import run as run_module


@pytest.fixture
def pipeline_stages(monkeypatch, tmp_path):
    """Keep orchestration real while substituting the expensive stage boundaries."""
    config = _config().model_copy(update={"enable_screenshot_replacement": False})
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    results = {
        "load_config": config,
        "_validate_pipeline_config": None,
        "_resolve_input_paths": {},
        "_build_missing_inputs_summary": {"missing_required": [], "optional_missing": []},
        "_validate_input_files": None,
        "_collect_source_report_dates": [],
        "_collect_cprs_header_candidates": [],
        "_create_run_directory": run_dir,
        "_prepare_runtime_config": config,
        "_parse_inputs": {},
        "_apply_daily_holdings_repo_cash": {},
        "_validate_parsed_inputs": None,
        "_run_reconciliation_checks": None,
        "_compute_metrics": ({}, {}),
        "_write_risk_outputs": [],
        "_compute_and_write_concentration_metrics": [],
        "_write_change_attribution_outputs": [],
        "_compute_and_write_limit_breaches": run_module.LimitBreachEvaluation(None, 0),
        "_update_historical_outputs": [],
        "_call_write_outputs": [],
        "_write_langsmith_fleet_artifact": run_dir / "fleet.json",
        "_resolve_manifest_input_paths": {},
    }
    stages = {}
    for name, result in results.items():
        stages[name] = Mock(return_value=result)
        monkeypatch.setattr(run_module, name, stages[name])
    builder = Mock()
    builder.build.return_value = {}
    builder.write.return_value = run_dir / "manifest.json"
    monkeypatch.setattr(run_module, "ManifestBuilder", Mock(return_value=builder))
    stages["manifest_build"] = builder.build
    stages["manifest_write"] = builder.write
    return stages, run_dir


@pytest.mark.parametrize(
    ("stage", "message"),
    [
        ("load_config", "config load"),
        ("_validate_pipeline_config", "config validation"),
        ("_validate_input_files", "input validation stage"),
        ("_collect_source_report_dates", "date derivation stage"),
        ("_create_run_directory", "run directory setup stage"),
        ("_prepare_runtime_config", "parse stage"),
        ("_compute_metrics", "compute stage"),
        ("_write_risk_outputs", "risk output stage"),
        ("_compute_and_write_concentration_metrics", "concentration metrics stage"),
        ("_write_change_attribution_outputs", "change attribution stage"),
        ("_compute_and_write_limit_breaches", "limit breach stage"),
        ("_update_historical_outputs", "historical update stage"),
        ("_call_write_outputs", "output write stage"),
        ("_write_langsmith_fleet_artifact", "LangSmith fleet artifact stage"),
        ("manifest_build", "manifest generation stage"),
        ("manifest_write", "manifest generation stage"),
    ],
)
def test_pipeline_stage_failures_preserve_cause_and_stop_later_stages(
    pipeline_stages, tmp_path, stage, message
) -> None:
    stages, _ = pipeline_stages
    failure = OSError("stage dependency failed")
    stages[stage].side_effect = failure
    with pytest.raises(RuntimeError, match=f"Pipeline failed during {message}") as raised:
        run_module.run_pipeline(tmp_path / "config.yml")
    assert raised.value.__cause__ is failure
    stages[stage].assert_called_once()
    ordered = list(stages)
    for later_stage in ordered[ordered.index(stage) + 1 :]:
        stages[later_stage].assert_not_called()


def test_run_pipeline_with_config_forwards_report_date_override(
    monkeypatch, tmp_path: Path
) -> None:
    """In-memory config execution must preserve an approved report-date override."""

    config_dir = tmp_path / "config"
    captured: dict[str, object] = {}

    def fake_run_pipeline(config_path: Path, **kwargs: object) -> Path:
        captured["config_path"] = config_path
        captured["kwargs"] = kwargs
        return tmp_path / "run"

    monkeypatch.setattr(run_module, "run_pipeline", fake_run_pipeline)
    result = run_module.run_pipeline_with_config(
        _config(),
        config_dir=config_dir,
        output_dir=tmp_path / "run",
        formatting_profile="accounting",
        report_date_override="source-date-confirmed",
    )

    assert result == tmp_path / "run"
    assert isinstance(captured["config_path"], Path)
    assert captured["config_path"].parent.name == "counter-risk-runtime"
    assert captured["kwargs"] == {
        "output_dir": tmp_path / "run",
        "formatting_profile": "accounting",
        "report_date_override": "source-date-confirmed",
    }
    assert not captured["config_path"].exists()


def test_fleet_artifact_excludes_external_output_paths_and_captures_environment(
    monkeypatch, tmp_path: Path
) -> None:
    """Fleet evidence references only artifacts contained in the run directory."""

    captured: dict[str, object] = {}

    def fake_build_fleet_records(**kwargs: object) -> list[dict[str, str]]:
        captured.update(kwargs)
        return [{"status": "ok"}]

    expected_path = tmp_path / "langsmith-fleet.json"
    monkeypatch.setattr(run_module, "build_fleet_records", fake_build_fleet_records)
    monkeypatch.setattr(run_module, "write_fleet_records", lambda path, records: expected_path)
    monkeypatch.setenv("LANGSMITH_TRACE_LATENCY_MS", " 123 ")
    monkeypatch.setenv("PIPELINE_ERROR_CATEGORY", " input ")

    result = run_module._write_langsmith_fleet_artifact(
        run_dir=tmp_path,
        as_of_date=date(2025, 12, 31),
        output_paths=[tmp_path / "manifest.json", tmp_path.parent / "outside.csv"],
        warnings=[],
        concentration_metrics_records=[{"counterparty": "A"}],
        risk_proxy_summary={"available": True},
        limit_breach_summary={"warning_breach_count": 1},
        workflow_trace_events=[{"stage": "parse"}],
    )

    assert result == expected_path
    assert captured["report_artifacts"] == ["artifact:manifest.json"]
    assert captured["risk_proxy_status"] == "success"
    assert captured["data_quality_status"] == "warning"
    assert captured["limit_max_severity"] == "warning"
    context = captured["context"]
    assert context.latency_ms == 123
    assert context.error_category == "input"


def test_limit_summary_and_failure_message_cover_invalid_and_fail_values(tmp_path: Path) -> None:
    """Limit policy must not turn malformed counts into a false failure."""

    no_breach = run_module.LimitBreachEvaluation(csv_path=None, breach_count=0)
    summary = run_module._build_limit_breach_summary(run_dir=tmp_path, limit_breaches=no_breach)
    assert summary["has_breaches"] is False
    assert summary["warning_banner"] is None
    run_module._raise_for_fail_limit_breaches(
        {"fail_breach_count": "not-a-number"}, run_dir=tmp_path
    )

    csv_path = tmp_path / "limit_breaches.csv"
    csv_path.write_text("severity\nfail\nfail\n", encoding="utf-8")
    banner = run_module._build_limit_warning_banner_for_run_dir(tmp_path)
    assert banner is not None
    assert banner.title == "Limit Breaches Detected (2)"

    fail_summary = {
        "fail_breach_count": 2,
        "report_path": "limit_breaches.csv",
    }
    try:
        run_module._raise_for_fail_limit_breaches(fail_summary, run_dir=tmp_path)
    except RuntimeError as exc:
        assert "Fail-severity limit breaches detected: 2" in str(exc)
        assert str(csv_path.resolve()) in str(exc)
    else:
        raise AssertionError("fail-severity limits must stop the pipeline")


@pytest.mark.parametrize(
    "field",
    [
        "mosers_all_programs_xlsx",
        "raw_nisa_all_programs_xlsx",
        "mosers_ex_trend_xlsx",
        "raw_nisa_ex_trend_xlsx",
        "mosers_trend_xlsx",
        "raw_nisa_trend_xlsx",
        "daily_holdings_pdf",
        "cash_overrides_csv",
        "dropin_all_programs_template_xlsx",
        "hist_all_programs_3yr_xlsx",
        "hist_ex_llc_3yr_xlsx",
        "hist_llc_3yr_xlsx",
        "monthly_pptx",
    ],
)
def test_pipeline_rejects_wrong_file_type_for_each_input(field, tmp_path) -> None:
    config = _config().model_copy(update={field: tmp_path / "wrong.txt"})
    with pytest.raises(ValueError, match=f"Invalid file type for {field}:"):
        run_module._validate_pipeline_config(config)


@pytest.mark.parametrize("variant", ["all_programs", "ex_trend", "trend"])
def test_pipeline_requires_a_generated_or_raw_source_for_each_variant(variant) -> None:
    config = _config().model_copy(update={f"mosers_{variant}_xlsx": None})
    with pytest.raises(ValueError, match=f"mosers_{variant}_xlsx.*must be configured"):
        run_module._validate_pipeline_config(config)


def test_pipeline_rejects_output_root_file_and_unsupported_cash_source(tmp_path) -> None:
    path = tmp_path / "existing.txt"
    path.write_text("occupied")
    with pytest.raises(ValueError, match="output_root must be a directory path"):
        run_module._validate_pipeline_config(_config().model_copy(update={"output_root": path}))
    with pytest.raises(ValueError, match="cash_source_path must be one of"):
        run_module._validate_pipeline_config(
            _config().model_copy(update={"cash_source_path": path})
        )
    assert path.read_text() == "occupied"


def test_optional_input_paths_are_validated_and_included_in_provenance(tmp_path) -> None:
    optional_paths = {
        "raw_nisa_all_programs_xlsx": tmp_path / "raw-all.XLSX",
        "raw_nisa_ex_trend_xlsx": tmp_path / "raw-ex.xlsx",
        "raw_nisa_trend_xlsx": tmp_path / "raw-trend.xlsx",
        "daily_holdings_pdf": tmp_path / "holdings.PDF",
        "cash_source_path": tmp_path / "cash.XLSM",
        "cash_overrides_csv": tmp_path / "overrides.CSV",
        "dropin_all_programs_template_xlsx": tmp_path / "dropin.xlsx",
    }
    config = _config().model_copy(update=optional_paths)
    run_module._validate_pipeline_config(config)
    paths = run_module._resolve_input_paths(config)
    assert {key: paths[key] for key in optional_paths} == optional_paths


@pytest.mark.parametrize("source_type", ["csv", "xlsx", "pdf"])
def test_explicit_cash_source_requires_a_path(source_type) -> None:
    config = _config().model_copy(update={"cash_source_type": source_type})
    with pytest.raises(ValueError, match=f"cash_source_type={source_type} requires"):
        run_module._resolve_repo_cash_source(config)


@pytest.mark.parametrize(
    ("suffix", "source_type"),
    [(".CSV", "csv"), (".xlsx", "xlsx"), (".xlsm", "xlsx"), (".pdf", "pdf")],
)
def test_cash_source_is_inferred_from_extension(suffix, source_type, tmp_path) -> None:
    path = tmp_path / f"cash{suffix}"
    config = _config().model_copy(update={"cash_source_path": path})
    assert run_module._resolve_repo_cash_source(config) == (source_type, path)


def test_cash_source_rejects_unknown_extension_and_falls_back_to_daily_pdf(tmp_path) -> None:
    config = _config().model_copy(update={"cash_source_path": tmp_path / "cash.txt"})
    with pytest.raises(ValueError, match="Unsupported cash_source_path extension: .txt"):
        run_module._resolve_repo_cash_source(config)
    pdf = tmp_path / "holdings.pdf"
    config = config.model_copy(update={"cash_source_path": None, "daily_holdings_pdf": pdf})
    assert run_module._resolve_repo_cash_source(config) == ("pdf", pdf)


@pytest.mark.parametrize("total_notional", [0.0, 100.0])
def test_historical_class_mix_falls_back_to_rollup_records(tmp_path, total_notional) -> None:
    from openpyxl import Workbook, load_workbook

    path = tmp_path / "history.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Class"
    # A subset of class headers is legal; missing classes must not shift columns.
    sheet.append(["Date", "Equity", "Treasury"])
    sheet.append([date(2025, 11, 30), 0.0, 1.0])
    book.save(path)
    book.close()
    run_module._merge_historical_workbook(
        workbook_path=path,
        variant="trend",
        as_of_date=date(2025, 12, 31),
        cprs_ch_records=[
            {"Segment": "totals", "Notional": total_notional, "Equity": 25, "Treasury": 75},
            {"Segment": "repo", "Notional": 999, "Equity": 999},
        ],
        warnings=[],
    )
    updated = load_workbook(path)
    try:
        row = list(updated["Class"].values)[-1]
        assert row[0].date() == date(2025, 12, 31)
        assert row[1:] == ((0.25, 0.75) if total_notional else (0, 0))
    finally:
        updated.close()


def test_historical_class_mix_prefers_published_absolute_values_and_defaults_missing_classes(
    tmp_path,
) -> None:
    from openpyxl import Workbook, load_workbook

    path = tmp_path / "history.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Class"
    sheet.append(["Date", "Equity", "Treasury"])
    sheet.append([date(2025, 11, 30), 0.25, 0.75])
    book.save(path)
    book.close()
    warnings = []

    run_module._merge_historical_workbook(
        workbook_path=path,
        variant="trend",
        as_of_date=date(2025, 12, 31),
        cprs_ch_records=[{"Segment": "totals", "Notional": 100, "Equity": -25, "Treasury": 125}],
        class_breakdown={"Equity": 0.8},
        warnings=warnings,
    )

    updated = load_workbook(path)
    try:
        rows = list(updated["Class"].values)
        assert len(rows) == 3
        assert rows[0] == ("Date", "Equity", "Treasury")
        assert rows[1][1:] == (0.25, 0.75)
        assert rows[2][0].date() == date(2025, 12, 31)
        assert rows[2][1:] == (0.8, 0)
    finally:
        updated.close()
    assert warnings == []


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (None, False),
        (False, False),
        (True, False),
        (0, False),
        (-2.5, True),
        (" #N/A ", False),
        ("—", False),
        (" 0.0 ", False),
        ("12", True),
    ],
)
def test_prior_historical_positions_ignore_blank_zero_and_error_cells(value, expected) -> None:
    assert run_module._historical_cell_has_value(value) is expected


def test_prior_populated_series_uses_last_dated_row_and_ignores_footer(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "history.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "Total"
    sheet.append(["Report title"])
    sheet.append(["Date", "Active", "Zero", "Error", "Dormant", "Boolean", ""])
    sheet.append([date(2025, 10, 31), 0, 10, 10, 10, 10, 10])
    sheet.append([None])
    sheet.append(["2025-11-30", -25, 0, "#N/A", None, True, 123])
    sheet.append(["Source: monthly report"])
    empty = book.create_sheet("Empty")
    empty.append(["As of Date", "Active"])
    notes = book.create_sheet("Notes")
    notes.append(["No date header"])
    book.save(path)
    book.close()

    assert run_module._extract_prior_populated_series_by_sheet(path) == {
        "Total": ("Active",),
        "Empty": (),
    }


@pytest.mark.parametrize("protocol", ["to_dict", "to_records"])
def test_table_record_adapters_filter_non_mapping_rows(protocol) -> None:
    def convert(**kwargs):
        assert kwargs == ({"orient": "records"} if protocol == "to_dict" else {})
        return [{"counterparty": "A", "Notional": 20}, None, "bad row"]

    table = SimpleNamespace(**{protocol: convert})
    assert run_module._records(table) == [{"counterparty": "A", "Notional": 20}]
    assert run_module._column_names(table) == {"counterparty", "Notional"}
    assert set(run_module._ordered_column_names(table)) == {"counterparty", "Notional"}


@pytest.mark.parametrize("columns", [None, 123])
def test_table_column_adapters_tolerate_non_iterable_columns(columns) -> None:
    table = SimpleNamespace(columns=columns)
    assert run_module._column_names(table) == set()
    assert run_module._ordered_column_names(table) == ()


@pytest.mark.parametrize(
    ("result", "error", "message"),
    [
        (None, RuntimeError, "did not record a refresh result"),
        (SimpleNamespace(), TypeError, "missing a status"),
        (SimpleNamespace(status="unknown"), ValueError, "Unsupported PPT processing status"),
    ],
)
def test_ppt_refresh_adapter_rejects_missing_or_unknown_status(result, error, message) -> None:
    with pytest.raises(error, match=message):
        run_module._to_ppt_processing_result(result)


@pytest.mark.parametrize("status", ["success", "failed"])
def test_ppt_refresh_adapter_normalizes_status_and_error_detail(status) -> None:
    result = run_module._to_ppt_processing_result(
        SimpleNamespace(status=SimpleNamespace(value=f" {status.upper()} "), error_detail=503)
    )
    assert result.status.value == status
    assert result.error_detail == "503"


def test_runtime_paths_resolve_screenshots_and_discovery_roots_without_mutating_config(
    tmp_path,
) -> None:
    absolute = tmp_path / "absolute.png"
    config = _config().model_copy(
        update={
            "screenshot_inputs": {"relative": Path("relative.png"), "absolute": absolute},
            "input_discovery": _config().input_discovery.model_copy(
                update={"directory_roots": {"relative": Path("inputs"), "absolute": tmp_path}}
            ),
        }
    )
    resolved = run_module._resolve_runtime_config_paths(config=config, config_dir=tmp_path)
    assert resolved.screenshot_inputs == {
        "relative": tmp_path / "relative.png",
        "absolute": absolute,
    }
    assert resolved.input_discovery.directory_roots == {
        "relative": tmp_path / "inputs",
        "absolute": tmp_path,
    }
    assert config.screenshot_inputs["relative"] == Path("relative.png")
    assert (
        run_module._resolve_runtime_config_paths(config=resolved, config_dir=tmp_path) == resolved
    )
    without_mappings = resolved.model_copy(
        update={"screenshot_inputs": {}, "input_discovery": _config().input_discovery}
    )
    assert (
        run_module._resolve_runtime_config_paths(config=without_mappings, config_dir=tmp_path)
        is without_mappings
    )


@pytest.mark.parametrize("payload", [None, {"totals": []}])
def test_parsed_input_validation_rejects_malformed_variant_payload(payload) -> None:
    parsed = dict.fromkeys(("all_programs", "ex_trend", "trend"), payload)
    with pytest.raises(ValueError, match="Parsed payload for variant 'all_programs'"):
        run_module._validate_parsed_inputs(parsed)


def test_required_columns_error_identifies_section_and_missing_columns() -> None:
    with pytest.raises(ValueError, match="trend.totals.*Notional, NotionalChange"):
        run_module._require_columns(
            section_name="trend.totals",
            columns={"counterparty"},
            required_columns=("counterparty", "Notional", "NotionalChange"),
        )


@pytest.mark.parametrize("cycle", ["self", "pair"])
def test_operator_diagnostics_terminate_on_exception_cycles(cycle) -> None:
    error = RuntimeError("outer")
    cause = error if cycle == "self" else RuntimeError("inner")
    error.__cause__ = cause
    cause.__cause__ = error
    assert run_module._find_unmapped_counterparty_error(error) is None
    assert run_module._build_reconciliation_failure_operator_message(error) == ""


@pytest.mark.parametrize(
    ("value", "expected"),
    [
        (datetime(2025, 12, 31, 15, 30), date(2025, 12, 31)),
        (" 12/31/25 ", date(2025, 12, 31)),
        ("2025/12/31", date(2025, 12, 31)),
        (" ", None),
        ("Report footer", None),
        (123, None),
    ],
)
def test_historical_row_date_recognizes_supported_dates_only(value, expected) -> None:
    assert run_module._coerce_historical_row_date(value) == expected


def test_unknown_historical_variant_warns_without_reading_or_creating_file(tmp_path) -> None:
    path = tmp_path / "missing.xlsx"
    warnings = []
    run_module._merge_historical_workbook(
        workbook_path=path,
        variant="unknown",
        as_of_date=date(2025, 12, 31),
        cprs_ch_records=[],
        warnings=warnings,
    )
    assert warnings == [
        "Historical workbook update skipped for variant 'unknown'; no sheet mapping"
    ]
    assert not path.exists()


def test_unrecognized_historical_sheets_fail_without_rewriting_workbook(tmp_path) -> None:
    from openpyxl import Workbook

    path = tmp_path / "unknown.xlsx"
    book = Workbook()
    book.active.title = "Notes"
    book.save(path)
    book.close()
    before = path.read_bytes()
    with pytest.raises(RuntimeError, match="Failed to update historical workbook") as raised:
        run_module._merge_historical_workbook(
            workbook_path=path,
            variant="trend",
            as_of_date=date(2025, 12, 31),
            cprs_ch_records=[],
            warnings=[],
        )
    assert isinstance(raised.value.__cause__, ValueError)
    assert "No recognized historical sheets" in str(raised.value.__cause__)
    assert path.read_bytes() == before


def test_worksheet_row_bound_can_fall_back_when_dimensions_are_unavailable() -> None:
    assert run_module._worksheet_max_row(SimpleNamespace(), fallback=25) == 25


@pytest.mark.parametrize("key", ["Slide x", "slide 2:", "2:no", "2:1:3"])
def test_zip_screenshot_keys_reject_invalid_slide_or_picture_numbers(key) -> None:
    with pytest.raises(ValueError, match="Invalid screenshot key for zip backend"):
        run_module._parse_zip_screenshot_key(key)


@pytest.mark.parametrize(("key", "expected"), [(" SLIDE 2 : 3 ", (2, 3)), ("4", (4, 0))])
def test_zip_screenshot_keys_preserve_slide_and_picture_numbers(key, expected) -> None:
    assert run_module._parse_zip_screenshot_key(key) == expected


@pytest.mark.parametrize("invalid_input", ["blank-key", "missing-file", "directory", "empty"])
def test_screenshot_mapping_rejects_unusable_inputs(tmp_path, invalid_input) -> None:
    image = tmp_path / "image.png"
    inputs = {"all programs": image}
    if invalid_input == "blank-key":
        image.write_bytes(b"unchanged image")
        inputs = {" \t ": image}
        error, message = ValueError, "Screenshot input keys must be non-empty"
    elif invalid_input == "empty":
        inputs = {}
        error, message = ValueError, (
            "Screenshot replacement is enabled but no screenshot_inputs were provided"
        )
    else:
        if invalid_input == "directory":
            image.mkdir()
        error, message = FileNotFoundError, (
            f"Screenshot input 'all programs' file does not exist: {image}"
        )
    config = _config().model_copy(
        update={"enable_screenshot_replacement": True, "screenshot_inputs": inputs}
    )
    with pytest.raises(error) as raised:
        run_module._resolve_screenshot_input_mapping(config)
    assert str(raised.value) == message
    if invalid_input == "blank-key":
        assert image.read_bytes() == b"unchanged image"


def test_screenshot_replacer_rejects_unknown_backend() -> None:
    with pytest.raises(
        ValueError, match="^Unsupported screenshot replacement implementation: typo$"
    ):
        run_module._get_screenshot_replacer("typo")


def test_python_pptx_screenshot_backend_replaces_image_and_preserves_source(tmp_path) -> None:
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    old_image = tmp_path / "old.png"
    new_image = tmp_path / "new.png"
    Image.new("RGB", (20, 20), "red").save(old_image)
    Image.new("RGB", (20, 20), "blue").save(new_image)
    source = tmp_path / "source.pptx"
    output = tmp_path / "output.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Custom section"
    picture = slide.shapes.add_picture(str(old_image), Inches(1), Inches(2), Inches(3))
    geometry = (picture.left, picture.top, picture.width, picture.height)
    deck.save(source)
    source_bytes = source.read_bytes()

    replacer = run_module._get_screenshot_replacer("python-pptx")
    replacer(source, output, {"Custom section": new_image})

    assert source.read_bytes() == source_bytes
    updated_slide = Presentation(output).slides[0]
    assert updated_slide.shapes.title.text == "Custom section"
    assert len(updated_slide.shapes) == 2
    updated_picture = updated_slide.shapes[1]
    assert updated_picture.image.blob == new_image.read_bytes()
    assert (
        updated_picture.left,
        updated_picture.top,
        updated_picture.width,
        updated_picture.height,
    ) == geometry


@pytest.mark.parametrize(
    ("require", "message"),
    [
        (run_module._require_parsed_by_variant, "parsed_by_variant is required"),
        (run_module._require_source_pptx, "source_pptx is required"),
    ],
)
def test_output_generation_requires_its_stage_inputs(require, message) -> None:
    with pytest.raises(ValueError, match=message):
        require(None)


@pytest.mark.parametrize(
    ("parsed", "message"),
    [
        (None, "parsed_by_variant is required"),
        ({}, "Parsed payload for variant 'all_programs' is missing$"),
        ({"all_programs": []}, "Parsed payload for variant 'all_programs' is missing$"),
        ({"all_programs": {}}, "missing the totals section required for Drop-In"),
    ],
    ids=["no-parsed-inputs", "missing-variant", "invalid-variant", "missing-totals"],
)
def test_dropin_rejects_incomplete_inputs_before_writing(
    monkeypatch, tmp_path, parsed, message
) -> None:
    config = _config().model_copy(
        update={"dropin_all_programs_template_xlsx": tmp_path / "template.xlsx"}
    )
    fill = Mock()
    monkeypatch.setattr(run_module, "fill_dropin_template", fill)
    warnings = []

    with pytest.raises(ValueError, match=message):
        run_module._write_all_programs_dropin_output(
            run_dir=tmp_path, config=config, warnings=warnings, parsed_by_variant=parsed
        )

    fill.assert_not_called()
    assert warnings == []
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("output_count", [0, 2])
def test_screenshot_generator_rejects_ambiguous_master_before_refresh(
    monkeypatch, tmp_path, output_count
) -> None:
    inputs = {}
    for variant in ("all_programs", "ex_trend", "trend"):
        source = tmp_path / f"{variant}-source.xlsx"
        source.write_bytes(b"input")
        inputs[f"mosers_{variant}_xlsx"] = source
    config = _config().model_copy(update=inputs)
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    generator = Mock(name="screenshot_generator")
    generator.name = "ppt_screenshot"
    generator.generate.return_value = tuple(
        run_dir / f"master-{index}.pptx" for index in range(output_count)
    )
    registry = Mock()
    registry.load.side_effect = [[generator], []]
    monkeypatch.setattr(run_module, "_build_output_generator_registry", lambda **kwargs: registry)
    refresh = Mock(side_effect=AssertionError("invalid Master must not reach Office"))
    monkeypatch.setattr(run_module, "_refresh_ppt_links", refresh)

    with pytest.raises(RuntimeError, match="must produce exactly one Master PPT output"):
        run_module._write_outputs(
            run_dir=run_dir, config=config, as_of_date=date(2025, 12, 31), warnings=[]
        )

    generator.generate.assert_called_once()
    assert registry.load.call_count == 1
    refresh.assert_not_called()
    assert not list(run_dir.glob("*.pptx"))


@pytest.mark.parametrize("auxiliary_output", [False, True], ids=["no-generators", "auxiliary-only"])
def test_missing_master_generator_preserves_workbooks_and_skips_ppt_stages(
    monkeypatch, tmp_path, auxiliary_output
) -> None:
    inputs = {}
    expected_workbooks = []
    for variant in ("all_programs", "ex_trend", "trend"):
        source = tmp_path / f"{variant}-source.xlsx"
        source.write_bytes(variant.encode())
        inputs[f"mosers_{variant}_xlsx"] = source
        expected_workbooks.append(tmp_path / "run" / f"{variant}-mosers-input.xlsx")
    config = _config().model_copy(update=inputs)
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    generator = Mock()
    generator.name = "auxiliary"
    auxiliary_path = run_dir / "auxiliary.csv"
    generator.generate.return_value = (auxiliary_path,)
    registry = Mock()
    registry.load.return_value = [generator] if auxiliary_output else []
    monkeypatch.setattr(run_module, "_build_output_generator_registry", lambda **kwargs: registry)
    refresh = Mock(side_effect=AssertionError("no Master exists to refresh"))
    derive = Mock(side_effect=AssertionError("no Master exists to distribute"))
    monkeypatch.setattr(run_module, "_refresh_ppt_links", refresh)
    monkeypatch.setattr(run_module, "_derive_distribution_ppt", derive)
    warnings = []

    paths, result = run_module._write_outputs(
        run_dir=run_dir, config=config, as_of_date=date(2025, 12, 31), warnings=warnings
    )

    assert paths == expected_workbooks + ([auxiliary_path] if auxiliary_output else [])
    assert [path.read_bytes() for path in expected_workbooks] == [
        variant.encode() for variant in ("all_programs", "ex_trend", "trend")
    ]
    assert result.status is run_module.PptProcessingStatus.SKIPPED
    assert result.ppt_outputs is None
    assert warnings == []
    registry.load.assert_called_once()
    assert registry.load.call_args.kwargs["stage"] == "ppt_master"
    if auxiliary_output:
        generator.generate.assert_called_once()
    else:
        generator.generate.assert_not_called()
    refresh.assert_not_called()
    derive.assert_not_called()
    assert not list(run_dir.glob("*.pptx"))
    assert not (run_dir / "README.txt").exists()


@pytest.mark.parametrize("detail", ["Office unavailable", ""], ids=["with-detail", "no-detail"])
def test_missing_refresh_generator_uses_fallback_and_reports_failure_without_distribution(
    monkeypatch, tmp_path, detail
) -> None:
    inputs = {}
    for variant in ("all_programs", "ex_trend", "trend"):
        source = tmp_path / f"{variant}-source.xlsx"
        source.write_bytes(variant.encode())
        inputs[f"mosers_{variant}_xlsx"] = source
    config = _config().model_copy(update=inputs)
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    master = run_dir / "master.pptx"
    master.write_bytes(b"master")
    generator = Mock()
    generator.name = "ppt_screenshot"
    generator.generate.return_value = (master,)
    registry = Mock()
    registry.load.side_effect = [[generator], []]
    monkeypatch.setattr(run_module, "_build_output_generator_registry", lambda **kwargs: registry)
    refresh = Mock(side_effect=RuntimeError(detail))
    derive = Mock(side_effect=AssertionError("failed Master must not be distributed"))
    monkeypatch.setattr(run_module, "_refresh_ppt_links", refresh)
    monkeypatch.setattr(run_module, "_derive_distribution_ppt", derive)
    warnings = []

    paths, result = run_module._write_outputs(
        run_dir=run_dir, config=config, as_of_date=date(2025, 12, 31), warnings=warnings
    )

    refresh.assert_called_once_with(master)
    derive.assert_not_called()
    assert [call.kwargs["stage"] for call in registry.load.call_args_list] == [
        "ppt_master",
        "ppt_refresh",
    ]
    assert warnings == [f"PPT links refresh failed; {detail or 'COM refresh encountered an error'}"]
    assert result.status is run_module.PptProcessingStatus.FAILED
    assert result.error_detail == detail
    assert result.ppt_outputs["master"]["status"] == "failed"
    assert result.ppt_outputs["distribution"]["status"] == "skipped"
    assert result.ppt_outputs["distribution"]["skipped_reason"] == (
        "Master PPT refresh failed; distribution PPT derivation skipped."
    )
    assert len(paths) == 4
    assert paths[-1] == master
    assert list(run_dir.glob("*.pptx")) == [master]
    assert not (run_dir / "README.txt").exists()


def test_pdf_export_rejects_multiple_outputs_instead_of_selecting_one(
    monkeypatch, tmp_path
) -> None:
    generator = Mock()
    generator.generate.return_value = (tmp_path / "first.pdf", tmp_path / "second.pdf")
    monkeypatch.setattr(
        run_module, "_build_pdf_export_output_generator", lambda **kwargs: generator
    )

    with pytest.raises(RuntimeError, match="must produce at most one PDF output"):
        run_module._export_distribution_pdf(
            source_pptx=tmp_path / "distribution.pptx",
            run_dir=tmp_path,
            config=_config(),
            as_of_date=date(2025, 12, 31),
            warnings=[],
        )

    generator.generate.assert_called_once()
    assert list(tmp_path.iterdir()) == []


def test_historical_helpers_report_unavailable_reader(monkeypatch, tmp_path) -> None:
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    path = tmp_path / "history.xlsx"
    warnings = []
    run_module._merge_historical_workbook(
        workbook_path=path,
        variant="trend",
        as_of_date=date(2025, 12, 31),
        cprs_ch_records=[],
        warnings=warnings,
    )
    assert warnings == [
        "Historical workbook update skipped for variant 'trend'; openpyxl unavailable"
    ]
    assert run_module._extract_prior_populated_series_by_sheet(path) is None
    with pytest.raises(RuntimeError, match="requires openpyxl") as raised:
        run_module._extract_historical_series_headers_by_sheet(path)
    assert isinstance(raised.value.__cause__, ImportError)
    assert not path.exists()


def test_prior_populated_series_falls_back_when_workbook_is_unreadable(tmp_path) -> None:
    assert run_module._extract_prior_populated_series_by_sheet(tmp_path / "missing.xlsx") is None


@pytest.mark.parametrize("failure_stage", ["lookup", "rows"])
def test_prior_populated_series_retains_readable_sheets_after_one_sheet_fails(
    tmp_path, monkeypatch, failure_stage
) -> None:
    """A damaged worksheet must not discard evidence from readable worksheets."""
    broken_sheet = Mock()
    broken_sheet.iter_rows.side_effect = ValueError("damaged worksheet")
    readable_sheet = Mock()
    readable_sheet.iter_rows.return_value = [
        ("Date", "Active", "Dormant"),
        (date(2025, 11, 30), 25, 0),
    ]

    def lookup(name):
        if name == "Broken":
            if failure_stage == "lookup":
                raise KeyError("missing worksheet")
            return broken_sheet
        assert name == "Total"
        return readable_sheet

    workbook = MagicMock(sheetnames=["Broken", "Total"])
    workbook.__getitem__.side_effect = lookup
    load = Mock(return_value=workbook)
    monkeypatch.setattr("openpyxl.load_workbook", load)
    path = tmp_path / "history.xlsx"

    assert run_module._extract_prior_populated_series_by_sheet(path) == {"Total": ("Active",)}
    load.assert_called_once_with(filename=path, read_only=True, data_only=True)
    readable_sheet.iter_rows.assert_called_once_with(values_only=True)
    workbook.close.assert_called_once_with()


@pytest.mark.parametrize("failure", [RuntimeError(), IndexError("broken shared string table")])
def test_historical_header_read_failure_preserves_context_cause_and_closes_workbook(
    tmp_path, monkeypatch, failure
) -> None:
    """Unexpected cell errors retain their type and location and release the workbook."""
    path = tmp_path / "history.xlsx"

    def read_cell(*, row, column):
        assert row == 1
        if column == 1:
            return SimpleNamespace(value="Date")
        raise failure

    worksheet = SimpleNamespace(title="Total", max_row=1, max_column=2, cell=read_cell)
    workbook = MagicMock(sheetnames=["Total"])
    workbook.__getitem__.return_value = worksheet
    load = Mock(return_value=workbook)
    monkeypatch.setattr("openpyxl.load_workbook", load)

    with pytest.raises(
        type(failure), match="Unexpected error while reading historical series headers"
    ) as raised:
        run_module._extract_historical_series_headers_by_sheet(path)

    assert raised.value.__cause__ is failure
    expected = f"Unexpected error while reading historical series headers in workbook {path}, sheet 'Total'"
    if str(failure):
        expected += f": {failure}"
    assert str(raised.value) == expected
    load.assert_called_once_with(filename=path, read_only=True, data_only=True)
    workbook.close.assert_called_once_with()


@pytest.mark.parametrize("error_type", [KeyError, ValueError, OSError])
def test_historical_header_open_failure_reports_path_and_preserves_cause(
    tmp_path, monkeypatch, error_type
) -> None:
    path = tmp_path / "history.xlsx"
    failure = error_type("unreadable workbook")
    load = Mock(side_effect=failure)
    monkeypatch.setattr("openpyxl.load_workbook", load)

    with pytest.raises(RuntimeError) as raised:
        run_module._extract_historical_series_headers_by_sheet(path)

    assert str(raised.value) == f"Unable to extract historical series headers from workbook: {path}"
    assert raised.value.__cause__ is failure
    load.assert_called_once_with(filename=path, read_only=True, data_only=True)
    assert not path.exists()


@pytest.mark.parametrize("failure_stage", ["column-count", "header-cell"])
def test_historical_headers_retain_readable_sheets_after_series_headers_fail(
    tmp_path, monkeypatch, failure_stage
) -> None:
    """A malformed series header must not discard later readable worksheets."""
    from openpyxl import Workbook

    readable_book = Workbook()
    readable_sheet = readable_book.active
    readable_sheet.append(["Date", " Active ", None, "Dormant"])
    broken_sheet = SimpleNamespace(max_row=1, max_column=2)

    def read_cell(*, row, column):
        assert row == 1
        if column == 1:
            return SimpleNamespace(value="Date")
        raise KeyError("unreadable series header")

    broken_sheet.cell = Mock(side_effect=read_cell)
    if failure_stage == "column-count":
        broken_sheet.max_column = "invalid"
    sheets = {"Broken": broken_sheet, "Total": readable_sheet}
    workbook = MagicMock(sheetnames=list(sheets))
    workbook.__getitem__.side_effect = sheets.__getitem__
    load = Mock(return_value=workbook)
    monkeypatch.setattr("openpyxl.load_workbook", load)
    path = tmp_path / "history.xlsx"

    try:
        result = run_module._extract_historical_series_headers_by_sheet(path)
    finally:
        readable_book.close()

    assert result == {"Total": ("Active", "Dormant")}
    assert [call.args[0] for call in workbook.__getitem__.call_args_list] == ["Broken", "Total"]
    assert broken_sheet.cell.call_count == (1 if failure_stage == "column-count" else 2)
    load.assert_called_once_with(filename=path, read_only=True, data_only=True)
    workbook.close.assert_called_once_with()


@pytest.mark.parametrize("failure_stage", ["export", "rebuild", "scrub"])
def test_static_distribution_failure_warns_and_stops_before_publishing(
    tmp_path, monkeypatch, failure_stage
) -> None:
    """A failed static deck must not be reported as a successful deliverable."""
    source = tmp_path / "master.pptx"
    source.write_bytes(b"original master deck")
    target = tmp_path / "distribution.pptx"
    images = [tmp_path / "slide-1.png"]
    stages = {
        "export": Mock(return_value=images),
        "rebuild": Mock(),
        "scrub": Mock(),
    }
    stages[failure_stage].side_effect = OSError("static stage unavailable")
    monkeypatch.setattr(run_module.platform, "system", lambda: "Windows")
    monkeypatch.setattr(run_module.importlib.util, "find_spec", lambda name: object())
    monkeypatch.setattr(run_module, "export_ppt_slides_as_png_via_com", stages["export"])
    monkeypatch.setattr(run_module, "_rebuild_pptx_from_slide_images", stages["rebuild"])
    monkeypatch.setattr(run_module, "scrub_external_relationships_from_pptx", stages["scrub"])
    warnings = ["earlier warning"]

    result = run_module._create_static_distribution(
        source_pptx=source,
        run_dir=tmp_path,
        config=_config().model_copy(update={"distribution_static": True}),
        warnings=warnings,
        output_path=target,
    )

    assert result == []
    assert warnings == [
        "earlier warning",
        "distribution_static generation failed: static stage unavailable",
    ]
    assert source.read_bytes() == b"original master deck"
    stages["export"].assert_called_once_with(
        source_pptx=source, slide_images_dir=tmp_path / "_distribution_slides"
    )
    if failure_stage == "export":
        stages["rebuild"].assert_not_called()
    else:
        stages["rebuild"].assert_called_once_with(
            source_pptx=source, output_path=target, slide_images=images
        )
    if failure_stage == "scrub":
        stages["scrub"].assert_called_once_with(target, scrubbed_pptx_path=target)
    else:
        stages["scrub"].assert_not_called()


@pytest.mark.parametrize("fail_policy", ["warn", "strict"])
@pytest.mark.parametrize("source_rows", ["", "CIBC,\n"], ids=["header-only", "missing-cash"])
def test_empty_repo_cash_source_obeys_policy_and_preserves_totals(
    tmp_path, fail_policy, source_rows
) -> None:
    """An empty configured source must be distinguished from a skipped cash import."""
    source = tmp_path / "cash.csv"
    source_text = f"counterparty,cash_value\n{source_rows}"
    source.write_text(source_text, encoding="utf-8")
    overrides = tmp_path / "overrides.csv"
    overrides.write_text("counterparty,cash_value,note\n", encoding="utf-8")
    config = _config()
    config = config.model_copy(
        update={
            "cash_source_type": "csv",
            "cash_source_path": source,
            "cash_overrides_csv": overrides,
            "reconciliation": config.reconciliation.model_copy(update={"fail_policy": fail_policy}),
        }
    )
    parsed = {"all_programs": {"totals": [{"Counterparty": "CIBC", "Cash": 25.0}]}}
    warnings = ["earlier warning"]
    message = f"Repo Cash source 'csv:{source}' did not yield any counterparty values."

    if fail_policy == "strict":
        with pytest.raises(ValueError) as raised:
            run_module._apply_daily_holdings_repo_cash(
                config=config, parsed_by_variant=parsed, warnings=warnings
            )
        assert str(raised.value) == message
    else:
        summary = run_module._apply_daily_holdings_repo_cash(
            config=config, parsed_by_variant=parsed, warnings=warnings
        )
        assert summary["reconciliation_findings"] == [
            {"code": "empty_cash_source", "severity": "warn", "message": message}
        ]
        assert summary["skipped_reason"] is None
        assert summary["applied_to_totals"] is False
        assert summary["counterparty_count"] == 0
        assert summary["total_cash"] == 0.0
        assert summary["overrides_path"] == str(overrides)
        assert summary["applied_override_count"] == 0
        assert summary["override_audit_rows"] == []

    assert warnings == [
        "earlier warning",
        f"Repo Cash source used: csv:{source}",
        f"Repo Cash overrides file present but empty: {overrides}",
    ] + ([message] if fail_policy == "warn" else [])
    assert parsed == {"all_programs": {"totals": [{"Counterparty": "CIBC", "Cash": 25.0}]}}
    assert source.read_text(encoding="utf-8") == source_text
    assert overrides.read_text(encoding="utf-8") == "counterparty,cash_value,note\n"


@pytest.mark.parametrize("table_type", ["records", "dataframe"])
def test_compute_metrics_keeps_empty_variants_and_continues(table_type) -> None:
    """Empty source totals stay explicit without hiding later variants' metrics."""
    totals = []
    if table_type == "dataframe":
        from pandas import DataFrame

        totals = DataFrame(columns=["counterparty", "Notional", "NotionalChange"])
    parsed = {
        "all_programs": {"totals": totals},
        "trend": {"totals": [{"counterparty": "Active", "Notional": -25, "NotionalChange": -5}]},
    }

    exposures, changes = run_module._compute_metrics(parsed)

    assert exposures["all_programs"] == []
    assert changes == {
        "all_programs": [],
        "trend": [{"counterparty": "Active", "notional_change": -5.0}],
    }
    assert [(row["counterparty"], row["notional"]) for row in exposures["trend"]] == [
        ("Active", -25.0)
    ]


@pytest.mark.parametrize("missing_row", [0, 1])
def test_compute_metrics_suppresses_incomplete_changes_but_keeps_exposures(missing_row) -> None:
    """A partial change column must not present missing data as a zero change."""
    rows = [
        {"counterparty": "Positive", "Notional": 10, "NotionalChange": 2},
        {"counterparty": "Negative", "Notional": -30, "NotionalChange": -7},
    ]
    del rows[missing_row]["NotionalChange"]
    parsed = {
        "all_programs": {"totals": rows},
        "trend": {"totals": [{"counterparty": "Active", "Notional": 40, "NotionalChange": 3}]},
    }

    exposures, changes = run_module._compute_metrics(parsed)

    assert changes == {
        "all_programs": [],
        "trend": [{"counterparty": "Active", "notional_change": 3.0}],
    }
    assert [(row["counterparty"], row["notional"]) for row in exposures["all_programs"]] == [
        ("Negative", -30.0),
        ("Positive", 10.0),
    ]
    assert [(row["counterparty"], row["notional"]) for row in exposures["trend"]] == [
        ("Active", 40.0)
    ]
    assert "NotionalChange" not in rows[missing_row]


def _config():
    from counter_risk.config import WorkflowConfig

    return WorkflowConfig(
        as_of_date="2025-12-31",
        mosers_all_programs_xlsx=Path("all.xlsx"),
        mosers_ex_trend_xlsx=Path("ex.xlsx"),
        mosers_trend_xlsx=Path("trend.xlsx"),
        hist_all_programs_3yr_xlsx=Path("hist-all.xlsx"),
        hist_ex_llc_3yr_xlsx=Path("hist-ex.xlsx"),
        hist_llc_3yr_xlsx=Path("hist-trend.xlsx"),
        monthly_pptx=Path("report.pptx"),
        output_root=Path("runs"),
    )
