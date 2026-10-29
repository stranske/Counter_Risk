"""Tests for XLSM artifact generation from the base template."""

from __future__ import annotations

import re
from datetime import UTC, date, datetime
from pathlib import Path
from zipfile import ZipFile

import pytest

from counter_risk.build import xlsm


def _write_legacy_template(path: Path) -> None:
    """Create the supported two-sheet template that predates Config inputs."""
    with (
        ZipFile("assets/templates/counter_risk_template.xlsm") as source,
        ZipFile(path, "w") as output,
    ):
        for info in source.infolist():
            if info.filename == "xl/worksheets/sheet3.xml":
                continue
            data = source.read(info.filename)
            if info.filename == "xl/workbook.xml":
                data = re.sub(rb'    <sheet name="Config"[^>]*/>\n', b"", data)
                data = re.sub(rb"  <definedNames>.*?</definedNames>\n", b"", data, flags=re.DOTALL)
            elif info.filename == "xl/_rels/workbook.xml.rels":
                data = re.sub(rb'  <Relationship Id="rId5"[^>]*/>\n', b"", data)
            elif info.filename == "[Content_Types].xml":
                data = re.sub(
                    rb'  <Override PartName="/xl/worksheets/sheet3.xml"[^>]*/>\n', b"", data
                )
            output.writestr(info, data)


@pytest.mark.parametrize("via_cli", [False, True])
def test_config_sheet_injection_preserves_legacy_workbook_and_is_idempotent(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], via_cli: bool
) -> None:
    import openpyxl

    path = tmp_path / "legacy.xlsm"
    _write_legacy_template(path)
    with ZipFile(path) as archive:
        original = {
            info.filename: (info, archive.read(info.filename)) for info in archive.infolist()
        }

    if via_cli:
        assert xlsm.main(["--template-path", str(path), "--inject-config-sheet"]) == 0
        assert capsys.readouterr().out == f"Config sheet injected into: {path}\n"
    else:
        xlsm.inject_config_sheet(path)

    workbook = openpyxl.load_workbook(path, keep_vba=True)
    try:
        assert workbook.sheetnames == ["Runner", "ControlData", "Config"]
        assert workbook["Config"]["A1"].value == "Counter Risk — File Path Configuration"
        expected_labels = [
            "MOSERS All Programs (.xlsx)",
            "MOSERS Ex-Trend (.xlsx)",
            "MOSERS Trend (.xlsx)",
            "Historical All Programs – 3yr (.xlsx)",
            "Historical Ex-LLC – 3yr (.xlsx)",
            "Historical LLC – 3yr (.xlsx)",
            "Monthly Report Template (.pptx)",
        ]
        for row, label in enumerate(expected_labels, start=4):
            assert workbook["Config"].cell(row, 1).value == label
            assert workbook["Config"].cell(row, 2).value is None
        assert {
            name: list(value.destinations) for name, value in workbook.defined_names.items()
        } == {
            "RunnerConfig_MosersAllPrograms": [("Config", "$B$4")],
            "RunnerConfig_MosersExTrend": [("Config", "$B$5")],
            "RunnerConfig_MosersTrend": [("Config", "$B$6")],
            "RunnerConfig_HistAllPrograms3yr": [("Config", "$B$7")],
            "RunnerConfig_HistExLlc3yr": [("Config", "$B$8")],
            "RunnerConfig_HistLlc3yr": [("Config", "$B$9")],
            "RunnerConfig_MonthlyPptx": [("Config", "$B$10")],
        }
    finally:
        workbook.close()

    modified_parts = {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml"}
    with ZipFile(path) as archive:
        assert set(archive.namelist()) == {*original, "xl/worksheets/sheet3.xml"}
        for name, (info, data) in original.items():
            if name not in modified_parts:
                assert archive.read(name) == data
            assert archive.getinfo(name).date_time == info.date_time
            assert archive.getinfo(name).compress_type == info.compress_type

    injected = path.read_bytes()
    xlsm.inject_config_sheet(path)
    assert path.read_bytes() == injected


def _read_zip_text(path: Path, member: str) -> str:
    with ZipFile(path) as workbook, workbook.open(member) as handle:
        return handle.read().decode("utf-8", errors="ignore")


def test_build_xlsm_artifact_copies_template_and_keeps_required_excel_parts(tmp_path: Path) -> None:
    output_path = tmp_path / "Runner.generated.xlsm"

    built = xlsm.build_xlsm_artifact(
        template_path=Path("assets/templates/counter_risk_template.xlsm"),
        output_path=output_path,
        as_of_date=date(2026, 1, 31),
        run_date=datetime(2026, 2, 26, 12, 0, tzinfo=UTC),
        version="1.2.3",
    )

    assert built == output_path
    assert built.is_file()
    with ZipFile(built) as workbook:
        members = set(workbook.namelist())
    assert "[Content_Types].xml" in members
    assert "xl/workbook.xml" in members
    assert "xl/worksheets/sheet1.xml" in members
    assert "xl/worksheets/sheet2.xml" in members
    assert "xl/vbaProject.bin" in members
    assert "docProps/core.xml" in members


def test_build_xlsm_artifact_keeps_runner_controls_and_month_selector(tmp_path: Path) -> None:
    output_path = tmp_path / "Runner.generated.xlsm"

    xlsm.build_xlsm_artifact(
        template_path=Path("assets/templates/counter_risk_template.xlsm"),
        output_path=output_path,
        as_of_date=date(2026, 1, 31),
        run_date=datetime(2026, 2, 26, 12, 0, tzinfo=UTC),
        version="1.2.3",
    )

    workbook_xml = _read_zip_text(output_path, "xl/workbook.xml")
    runner_sheet_xml = _read_zip_text(output_path, "xl/worksheets/sheet1.xml")
    control_data_xml = _read_zip_text(output_path, "xl/worksheets/sheet2.xml")

    assert 'sheet name="Runner"' in workbook_xml
    assert 'sheet name="ControlData"' in workbook_xml
    assert "Run All" in runner_sheet_xml
    assert "Run Ex Trend" in runner_sheet_xml
    assert "Run Trend" in runner_sheet_xml
    assert "Open Output Folder" in runner_sheet_xml
    assert 'sqref="B3"' in runner_sheet_xml
    assert "ControlData!$A$2:$A$193" in runner_sheet_xml
    assert "MonthEnd" in control_data_xml


def test_build_xlsm_artifact_injects_metadata_into_core_properties(tmp_path: Path) -> None:
    output_path = tmp_path / "Runner.generated.xlsm"

    xlsm.build_xlsm_artifact(
        template_path=Path("assets/templates/counter_risk_template.xlsm"),
        output_path=output_path,
        as_of_date=date(2026, 1, 31),
        run_date=datetime(2026, 2, 26, 12, 34, 56, tzinfo=UTC),
        version="9.9.9",
    )

    core_xml = _read_zip_text(output_path, "docProps/core.xml")
    assert "Counter Risk Runner" in core_xml
    assert "as_of_date=2026-01-31" in core_xml
    assert "version=9.9.9" in core_xml
    assert "2026-02-26T12:34:56Z" in core_xml


def test_build_xlsm_artifact_retains_runnerlaunch_vba_markers(tmp_path: Path) -> None:
    output_path = tmp_path / "Runner.generated.xlsm"

    xlsm.build_xlsm_artifact(
        template_path=Path("assets/templates/counter_risk_template.xlsm"),
        output_path=output_path,
        as_of_date=date(2026, 1, 31),
        run_date=datetime(2026, 2, 26, 12, 0, tzinfo=UTC),
        version="1.2.3",
    )

    with ZipFile(output_path) as workbook, workbook.open("xl/vbaProject.bin") as handle:
        vba_text = handle.read().decode("latin-1", errors="ignore")

    assert 'Attribute VB_Name = "RunnerLaunch"' in vba_text
    assert "Public Sub RunAll_Click()" in vba_text
    assert "Public Sub RunExTrend_Click()" in vba_text
    assert "Public Sub RunTrend_Click()" in vba_text


def test_build_xlsm_artifact_raises_for_missing_template(tmp_path: Path) -> None:
    with pytest.raises(FileNotFoundError, match="Template XLSM was not found"):
        xlsm.build_xlsm_artifact(
            template_path=tmp_path / "missing-template.xlsm",
            output_path=tmp_path / "out.xlsm",
            as_of_date=date(2026, 1, 31),
            run_date=datetime(2026, 2, 26, 12, 0, tzinfo=UTC),
            version="1.2.3",
        )


def test_main_builds_xlsm_artifact_from_cli(tmp_path: Path) -> None:
    output_path = tmp_path / "Runner.generated.xlsm"

    result = xlsm.main(
        [
            "--template-path",
            "assets/templates/counter_risk_template.xlsm",
            "--output-path",
            str(output_path),
            "--as-of-date",
            "2026-01-31",
            "--run-date",
            "2026-02-26T12:00:00+00:00",
            "--version",
            "2.0.0",
        ]
    )

    assert result == 0
    assert output_path.is_file()


def test_build_xlsm_artifact_preserves_vba_project_across_version_updates(tmp_path: Path) -> None:
    output_v1 = tmp_path / "Runner.v1.xlsm"
    output_v2 = tmp_path / "Runner.v2.xlsm"

    xlsm.build_xlsm_artifact(
        template_path=Path("assets/templates/counter_risk_template.xlsm"),
        output_path=output_v1,
        as_of_date=date(2026, 1, 31),
        run_date=datetime(2026, 2, 26, 12, 0, tzinfo=UTC),
        version="1.2.3",
    )
    xlsm.build_xlsm_artifact(
        template_path=Path("assets/templates/counter_risk_template.xlsm"),
        output_path=output_v2,
        as_of_date=date(2026, 1, 31),
        run_date=datetime(2026, 2, 26, 12, 0, tzinfo=UTC),
        version="1.2.4",
    )

    with (
        ZipFile(output_v1) as workbook_v1,
        ZipFile(output_v2) as workbook_v2,
        workbook_v1.open("xl/vbaProject.bin") as vba_v1,
        workbook_v2.open("xl/vbaProject.bin") as vba_v2,
    ):
        assert vba_v1.read() == vba_v2.read()

    runner_sheet_v2_xml = _read_zip_text(output_v2, "xl/worksheets/sheet1.xml")
    assert "Run All" in runner_sheet_v2_xml
    assert "Run Ex Trend" in runner_sheet_v2_xml
    assert "Run Trend" in runner_sheet_v2_xml
    assert "Open Output Folder" in runner_sheet_v2_xml


def test_runner_xlsm_manual_macro_verification_doc_is_present_and_actionable() -> None:
    doc_path = Path("docs/runner_xlsm_macro_manual_verification.md")
    content = doc_path.read_text(encoding="utf-8")

    assert doc_path.is_file()
    assert "Manual Macro/Button Check" in content
    assert "Enable Content" in content
    assert "RunAll_Click" in content
    assert "RunExTrend_Click" in content
    assert "RunTrend_Click" in content
    assert "OpenOutputFolder_Click" in content
    assert "Version Bump Regression Check" in content
    assert "--version 1.2.4" in content
