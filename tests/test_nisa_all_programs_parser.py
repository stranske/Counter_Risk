"""Robustness tests for raw NISA All Programs parser."""

from __future__ import annotations

from dataclasses import asdict
from pathlib import Path
from typing import Any
from unittest.mock import Mock
from zipfile import ZipFile

import pytest

from counter_risk.parsers.nisa_all_programs import (
    NisaAllProgramsData,
    get_nisa_all_programs_input_structure,
    parse_nisa_all_programs,
)

_CANONICAL_HEADERS: tuple[str, ...] = (
    "counterparty",
    "cash",
    "tips",
    "treasury",
    "equity",
    "commodity",
    "currency",
    "notional",
    "notional_change",
    "annualized_volatility",
)

_HEADER_LABELS = {
    "counterparty": "Counterparty/ Clearing House",
    "cash": "Cash",
    "tips": "TIPS",
    "treasury": "Treasury",
    "equity": "Equity",
    "commodity": "Commodity",
    "currency": "Currency",
    "notional": "Notional",
    "notional_change": "from prior month***",
    "annualized_volatility": "%",
}


def _set_cell_if_header(
    worksheet: Any,
    *,
    row: int,
    header_columns: dict[str, int],
    header_name: str,
    value: object,
) -> None:
    if header_name in header_columns:
        worksheet.cell(row=row, column=header_columns[header_name]).value = value


def _build_raw_nisa_workbook(
    path: Path,
    *,
    sheet_name: str = "Data",
    include_leading_rows: int = 0,
    header_order: tuple[str, ...] = _CANONICAL_HEADERS,
) -> Path:
    openpyxl = pytest.importorskip("openpyxl")

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name

    header_row = 4 + include_leading_rows
    for row_number in range(1, header_row - 1):
        worksheet.cell(row=row_number, column=1).value = f"lead-in-{row_number}"

    header_columns: dict[str, int] = {}
    for offset, header_name in enumerate(header_order, start=2):
        header_columns[header_name] = offset
        worksheet.cell(row=header_row, column=offset).value = _HEADER_LABELS[header_name]
    worksheet.cell(row=header_row - 1, column=header_columns["annualized_volatility"]).value = (
        "Annualized Volatility"
    )

    segment_column = 1
    first_ch_row = header_row + 2
    worksheet.cell(row=first_ch_row, column=segment_column).value = "Swaps"
    _set_cell_if_header(
        worksheet,
        row=first_ch_row,
        header_columns=header_columns,
        header_name="counterparty",
        value="Alpha Bank",
    )
    _set_cell_if_header(
        worksheet, row=first_ch_row, header_columns=header_columns, header_name="cash", value=10.0
    )
    _set_cell_if_header(
        worksheet, row=first_ch_row, header_columns=header_columns, header_name="tips", value=20.0
    )
    _set_cell_if_header(
        worksheet,
        row=first_ch_row,
        header_columns=header_columns,
        header_name="treasury",
        value=30.0,
    )
    _set_cell_if_header(
        worksheet, row=first_ch_row, header_columns=header_columns, header_name="equity", value=40.0
    )
    _set_cell_if_header(
        worksheet,
        row=first_ch_row,
        header_columns=header_columns,
        header_name="commodity",
        value=50.0,
    )
    _set_cell_if_header(
        worksheet,
        row=first_ch_row,
        header_columns=header_columns,
        header_name="currency",
        value=60.0,
    )
    _set_cell_if_header(
        worksheet,
        row=first_ch_row,
        header_columns=header_columns,
        header_name="notional",
        value=210.0,
    )
    _set_cell_if_header(
        worksheet,
        row=first_ch_row,
        header_columns=header_columns,
        header_name="notional_change",
        value=5.0,
    )
    _set_cell_if_header(
        worksheet,
        row=first_ch_row,
        header_columns=header_columns,
        header_name="annualized_volatility",
        value=0.11,
    )

    second_ch_row = first_ch_row + 1
    _set_cell_if_header(
        worksheet,
        row=second_ch_row,
        header_columns=header_columns,
        header_name="counterparty",
        value="Beta Bank",
    )
    _set_cell_if_header(
        worksheet, row=second_ch_row, header_columns=header_columns, header_name="cash", value=1.0
    )
    _set_cell_if_header(
        worksheet, row=second_ch_row, header_columns=header_columns, header_name="tips", value=2.0
    )
    _set_cell_if_header(
        worksheet,
        row=second_ch_row,
        header_columns=header_columns,
        header_name="treasury",
        value=3.0,
    )
    _set_cell_if_header(
        worksheet, row=second_ch_row, header_columns=header_columns, header_name="equity", value=4.0
    )
    _set_cell_if_header(
        worksheet,
        row=second_ch_row,
        header_columns=header_columns,
        header_name="commodity",
        value=5.0,
    )
    _set_cell_if_header(
        worksheet,
        row=second_ch_row,
        header_columns=header_columns,
        header_name="currency",
        value=6.0,
    )
    _set_cell_if_header(
        worksheet,
        row=second_ch_row,
        header_columns=header_columns,
        header_name="notional",
        value=21.0,
    )
    _set_cell_if_header(
        worksheet,
        row=second_ch_row,
        header_columns=header_columns,
        header_name="notional_change",
        value=1.0,
    )
    _set_cell_if_header(
        worksheet,
        row=second_ch_row,
        header_columns=header_columns,
        header_name="annualized_volatility",
        value=0.2,
    )

    totals_marker_row = header_row + 20
    worksheet.cell(
        row=totals_marker_row,
        column=header_columns["counterparty"],
    ).value = "Total by Counterparty/Clearing House (This is not the legal obligation exposure)"
    _set_cell_if_header(
        worksheet,
        row=totals_marker_row + 1,
        header_columns=header_columns,
        header_name="counterparty",
        value="Alpha Bank",
    )
    _set_cell_if_header(
        worksheet,
        row=totals_marker_row + 1,
        header_columns=header_columns,
        header_name="tips",
        value=20.0,
    )
    _set_cell_if_header(
        worksheet,
        row=totals_marker_row + 1,
        header_columns=header_columns,
        header_name="treasury",
        value=30.0,
    )
    _set_cell_if_header(
        worksheet,
        row=totals_marker_row + 1,
        header_columns=header_columns,
        header_name="equity",
        value=40.0,
    )
    _set_cell_if_header(
        worksheet,
        row=totals_marker_row + 1,
        header_columns=header_columns,
        header_name="commodity",
        value=50.0,
    )
    _set_cell_if_header(
        worksheet,
        row=totals_marker_row + 1,
        header_columns=header_columns,
        header_name="currency",
        value=60.0,
    )
    _set_cell_if_header(
        worksheet,
        row=totals_marker_row + 1,
        header_columns=header_columns,
        header_name="notional",
        value=210.0,
    )
    _set_cell_if_header(
        worksheet,
        row=totals_marker_row + 1,
        header_columns=header_columns,
        header_name="notional_change",
        value=5.0,
    )
    _set_cell_if_header(
        worksheet,
        row=totals_marker_row + 1,
        header_columns=header_columns,
        header_name="annualized_volatility",
        value=0.11,
    )
    worksheet.cell(row=totals_marker_row + 2, column=header_columns["counterparty"]).value = (
        "Total Current Exposure"
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()
    return path


def _assert_expected_rows(data: NisaAllProgramsData) -> None:
    assert len(data.ch_rows) == 2
    assert data.ch_rows[0].segment == "swaps"
    assert data.ch_rows[0].counterparty == "Alpha Bank"
    assert data.ch_rows[0].cash == pytest.approx(10.0)
    assert data.ch_rows[0].notional == pytest.approx(210.0)
    assert len(data.totals_rows) == 1
    assert data.totals_rows[0].counterparty == "Alpha Bank"
    assert data.totals_rows[0].notional == pytest.approx(210.0)


def test_parse_nisa_all_programs_uses_fixture_file() -> None:
    fixture_path = Path("tests/fixtures/raw_nisa_all_programs.xlsx")
    assert fixture_path.exists(), f"Missing required fixture: {fixture_path}"

    parsed = parse_nisa_all_programs(fixture_path)
    _assert_expected_rows(parsed)


def test_parser_handles_when_data_worksheet_is_not_first_sheet(tmp_path: Path) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    workbook_path = tmp_path / "raw_nisa_non_first_sheet.xlsx"
    data_sheet_name = "NISA Data"
    _build_raw_nisa_workbook(workbook_path, sheet_name=data_sheet_name)

    workbook = openpyxl.load_workbook(workbook_path)
    try:
        first_sheet = workbook.create_sheet("Intro", 0)
        first_sheet.cell(row=1, column=1).value = "Readme"
        workbook.save(workbook_path)
    finally:
        workbook.close()

    parsed = parse_nisa_all_programs(workbook_path)
    _assert_expected_rows(parsed)


def test_parser_handles_extra_leading_rows_before_header_row(tmp_path: Path) -> None:
    workbook_path = tmp_path / "raw_nisa_leading_rows.xlsx"
    _build_raw_nisa_workbook(workbook_path, include_leading_rows=6)

    parsed = parse_nisa_all_programs(workbook_path)
    _assert_expected_rows(parsed)


@pytest.mark.parametrize("reverse_sheet_order", [False, True], ids=["original", "reversed"])
@pytest.mark.parametrize(
    ("first_name", "leading_rows", "second_name", "expected_name"),
    [("Alpha", 6, "Zeta", "Zeta"), ("Beta", 0, "alpha", "alpha")],
    ids=["earliest-header", "case-insensitive-name"],
)
def test_parser_selects_valid_sheet_deterministically(
    tmp_path: Path,
    first_name: str,
    leading_rows: int,
    second_name: str,
    expected_name: str,
    reverse_sheet_order: bool,
) -> None:
    import openpyxl

    path = _build_raw_nisa_workbook(
        tmp_path / "multiple-data-sheets.xlsx",
        sheet_name=first_name,
        include_leading_rows=leading_rows,
    )
    workbook = openpyxl.load_workbook(path)
    try:
        first = workbook.active
        second = workbook.copy_worksheet(first)
        second.title = second_name
        if leading_rows:
            second.delete_rows(1, leading_rows)
        # Distinct payloads expose which valid sheet supplied each parsed section.
        for worksheet, offset in [(first, leading_rows), (second, 0)]:
            worksheet.cell(6 + offset, 2).value = worksheet.title
            worksheet.cell(25 + offset, 2).value = worksheet.title
        if reverse_sheet_order:
            workbook.move_sheet(second, offset=-1)
        workbook.save(path)
    finally:
        workbook.close()

    parsed = parse_nisa_all_programs(path)

    assert [row.counterparty for row in parsed.ch_rows] == [expected_name, "Beta Bank"]
    assert [row.counterparty for row in parsed.totals_rows] == [expected_name]
    assert [row.notional for row in parsed.ch_rows] == pytest.approx([210, 21])
    assert parsed.totals_rows[0].notional == pytest.approx(210)


def test_parser_handles_reordered_columns(tmp_path: Path) -> None:
    workbook_path = tmp_path / "raw_nisa_reordered_columns.xlsx"
    header_order = (
        "tips",
        "counterparty",
        "notional_change",
        "currency",
        "cash",
        "equity",
        "notional",
        "commodity",
        "treasury",
        "annualized_volatility",
    )
    _build_raw_nisa_workbook(workbook_path, header_order=header_order)

    parsed = parse_nisa_all_programs(workbook_path)
    _assert_expected_rows(parsed)


@pytest.mark.parametrize(
    "header_order",
    [_CANONICAL_HEADERS, tuple(reversed(_CANONICAL_HEADERS))],
    ids=["standard-columns", "reversed-columns"],
)
def test_parser_preserves_each_numeric_field_in_ch_and_totals(
    tmp_path: Path, header_order: tuple[str, ...]
) -> None:
    import openpyxl

    path = _build_raw_nisa_workbook(tmp_path / "numeric-fields.xlsx", header_order=header_order)
    ch_values = {
        "cash": 101.25,
        "tips": 202.5,
        "treasury": 303.75,
        "equity": 404.25,
        "commodity": 505.5,
        "currency": 606.75,
        "notional": 707.25,
        "notional_change": -808.5,
        "annualized_volatility": 0.125,
    }
    totals_values = {
        "tips": -1202.5,
        "treasury": 1303.75,
        "equity": 1404.25,
        "commodity": 1505.5,
        "currency": 1606.75,
        "notional": 1707.25,
        "notional_change": -1808.5,
        "annualized_volatility": 0.25,
    }
    workbook = openpyxl.load_workbook(path)
    try:
        for column, header in enumerate(header_order, start=2):
            if header in ch_values:
                workbook.active.cell(6, column).value = ch_values[header]
            if header in totals_values:
                value = totals_values[header]
                # Totals exported by Excel can use accounting parentheses and grouping.
                workbook.active.cell(25, column).value = (
                    f"({-value:,.3f})" if value < 0 else f"{value:,.3f}"
                )
        workbook.save(path)
    finally:
        workbook.close()

    parsed = parse_nisa_all_programs(path)

    assert asdict(parsed.ch_rows[0]) == {
        "segment": "swaps",
        "counterparty": "Alpha Bank",
        **ch_values,
    }
    assert asdict(parsed.totals_rows[0]) == {"counterparty": "Alpha Bank", **totals_values}


def test_parser_raises_explicit_missing_headers_error(tmp_path: Path) -> None:
    workbook_path = tmp_path / "raw_nisa_missing_headers.xlsx"
    _build_raw_nisa_workbook(
        workbook_path,
        header_order=(
            "counterparty",
            "cash",
            "tips",
            "treasury",
            "equity",
            "notional",
            "notional_change",
            "annualized_volatility",
        ),
    )

    with pytest.raises(ValueError, match="Missing required headers:") as exc_info:
        parse_nisa_all_programs(workbook_path)

    message = str(exc_info.value)
    assert "commodity" in message
    assert "currency" in message


def test_parser_rejects_missing_and_non_workbook_paths(tmp_path: Path) -> None:
    missing_path = tmp_path / "missing.xlsx"
    with pytest.raises(FileNotFoundError, match="NISA raw workbook not found"):
        parse_nisa_all_programs(missing_path)

    text_path = tmp_path / "not-a-workbook.csv"
    text_path.write_text("counterparty", encoding="utf-8")
    with pytest.raises(ValueError, match="must be an .xlsx file"):
        parse_nisa_all_programs(text_path)


def test_parser_rejects_unreadable_workbook(tmp_path: Path) -> None:
    workbook_path = tmp_path / "unreadable.xlsx"
    workbook_path.write_text("not an xlsx archive", encoding="utf-8")

    with pytest.raises(ValueError, match="Unable to open NISA workbook"):
        parse_nisa_all_programs(workbook_path)


def test_parser_skips_subtotals_and_stops_totals_at_next_section(tmp_path: Path) -> None:
    openpyxl = pytest.importorskip("openpyxl")
    workbook_path = tmp_path / "raw_nisa_with_subtotals.xlsx"
    _build_raw_nisa_workbook(workbook_path)

    workbook = openpyxl.load_workbook(workbook_path)
    try:
        worksheet = workbook.active
        worksheet.cell(row=8, column=2).value = "Subtotal"
        worksheet.cell(row=26, column=2).value = "Total"
        worksheet.cell(row=27, column=2).value = "Total Current Exposure"
        worksheet.cell(row=28, column=2).value = "Should Not Be Parsed"
        workbook.save(workbook_path)
    finally:
        workbook.close()

    parsed = parse_nisa_all_programs(workbook_path)

    assert [row.counterparty for row in parsed.ch_rows] == ["Alpha Bank", "Beta Bank"]
    assert [row.counterparty for row in parsed.totals_rows] == ["Alpha Bank"]


@pytest.mark.parametrize(
    ("cleared_cells", "message"),
    [
        ([(24, 2)], "Unable to locate 'Total by Counterparty/Clearing House' section"),
        ([(6, 1)], "produced no CPRS-CH rows"),
        ([(25, 2)], "produced no totals rows"),
    ],
    ids=["missing-totals-marker", "missing-segment", "empty-totals"],
)
def test_parser_rejects_incomplete_sections(
    tmp_path: Path, cleared_cells: list[tuple[int, int]], message: str
) -> None:
    import openpyxl

    path = _build_raw_nisa_workbook(tmp_path / "incomplete.xlsx")
    workbook = openpyxl.load_workbook(path)
    try:
        for row, column in cleared_cells:
            workbook.active.cell(row, column).value = None
        workbook.save(path)
    finally:
        workbook.close()

    with pytest.raises(ValueError, match=message):
        parse_nisa_all_programs(path)


def test_parser_closes_workbook_with_no_worksheets(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "empty.xlsx"
    path.touch()
    workbook = Mock(sheetnames=[])
    monkeypatch.setattr("openpyxl.load_workbook", lambda **kwargs: workbook)

    with pytest.raises(ValueError, match="contains no worksheets"):
        parse_nisa_all_programs(path)

    workbook.close.assert_called_once_with()


def test_parser_infers_unlabelled_counterparty_column_from_totals_marker(tmp_path: Path) -> None:
    import openpyxl

    path = _build_raw_nisa_workbook(tmp_path / "unlabelled.xlsx")
    workbook = openpyxl.load_workbook(path)
    try:
        workbook.active.cell(4, 2).value = None
        # A label before the first segment must not become a counterparty row.
        workbook.active.cell(5, 2).value = "Preparation notes"
        workbook.active.cell(28, 3).value = 999
        workbook.save(path)
    finally:
        workbook.close()

    _assert_expected_rows(parse_nisa_all_programs(path))


def test_parser_reports_missing_totals_with_unlabelled_counterparty_column(tmp_path: Path) -> None:
    import openpyxl

    path = _build_raw_nisa_workbook(tmp_path / "no-marker.xlsx")
    workbook = openpyxl.load_workbook(path)
    try:
        workbook.active.cell(4, 2).value = None
        workbook.active.cell(24, 2).value = None
        workbook.save(path)
    finally:
        workbook.close()

    with pytest.raises(ValueError, match="Unable to locate 'Total by Counterparty"):
        parse_nisa_all_programs(path)


def test_parser_reads_workbook_without_dimension_record(tmp_path: Path) -> None:
    import re

    original = _build_raw_nisa_workbook(tmp_path / "original.xlsx")
    path = tmp_path / "dimensionless.xlsx"
    with ZipFile(original) as source, ZipFile(path, "w") as destination:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data, count = re.subn(rb"<dimension\b[^>]*/>", b"", data)
                assert count == 1
            destination.writestr(entry, data)

    _assert_expected_rows(parse_nisa_all_programs(path))


@pytest.mark.parametrize(
    "stop_marker", ["total current exposure", "mosers program", "notional breakdown"]
)
def test_published_input_contract_describes_parseable_workbook(
    tmp_path: Path, stop_marker: str
) -> None:
    import openpyxl

    structure = get_nisa_all_programs_input_structure()
    assert set(structure.required_headers) == set(_CANONICAL_HEADERS) - {"counterparty"}
    assert stop_marker in structure.totals_stop_markers
    assert structure.segment_aliases["futures / cdx"] == "futures_cdx"

    path = _build_raw_nisa_workbook(tmp_path / "published-contract.xlsx")
    workbook = openpyxl.load_workbook(path)
    try:
        worksheet = workbook.active
        for column, header in enumerate(_CANONICAL_HEADERS, start=2):
            worksheet.cell(4, column).value = structure.header_aliases[header][0]
        worksheet.cell(6, 1).value = "Futures / CDX"
        worksheet.cell(24, 2).value = structure.totals_marker
        worksheet.cell(26, 2).value = stop_marker.title()
        worksheet.cell(27, 2).value = "Later Section Bank"
        worksheet.cell(27, 9).value = 999
        workbook.save(path)
    finally:
        workbook.close()

    parsed = parse_nisa_all_programs(path)
    assert [row.segment for row in parsed.ch_rows] == ["futures_cdx", "futures_cdx"]
    assert [row.notional for row in parsed.ch_rows] == pytest.approx([210, 21])
    assert [row.counterparty for row in parsed.totals_rows] == ["Alpha Bank"]
    assert parsed.totals_rows[0].notional == pytest.approx(210)
